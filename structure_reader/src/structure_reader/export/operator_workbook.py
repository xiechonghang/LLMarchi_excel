"""HF decoder 阅读数据 → 两表 Excel 工作簿导出器（里程碑一，pi 侧独立实现）。

公开接口（由 structure_reader/docs/hf_reading_contract_v1.md 冻结）：

    export_operator_workbook(reading_document: dict) -> bytes

职责边界：
- 只呈现输入文档中的事实：不读取 models/、outputs/、.work/，不导入 HF 提取器、
  torch 或 Claude 的任何模块，不自己写文件，只返回字节。
- 输入不合法（缺字段、枚举越界、引用无法解析、id 重复、阶段数量/顺序错误）时
  抛 ReadingDocumentError（ValueError 子类），不静默补造缺失事实。
- 只生成 Prefill、Decode 两张表；每张表自成一体：阶段简述 → 参数符号区 → 算子区。
  没有第三张 sheet、没有隐藏 metadata sheet、没有重复次数列、不逐专家复制行。
- 参数引用落到本表参数区单元格：算子行的联查列是 Excel 公式，指向本表参数行的
  符号单元格与值单元格（公式栏可见引用，「参数引用」列可点击跳转），不是复制一份
  无法追溯的数字。shape 一律保留符号，不为了模拟运行填入 B/T/P 的固定值。
- 轴符号→参数行的绑定用本行显式 parameter_refs（算子及其输入/输出/权重张量）分辨：
  同名 symbol 候选只有一个时直接绑定；多个时只在显式引用能唯一分辨时绑定，否则明写
  「符号存在歧义」并列出候选参数id与本表行号，绝不默认取第一个、也不悄悄选错数值。
- analysis_status == internal_unanalyzed 的算子醒目标注中文「内部未分析」。
- 相同输入产出相同字节：固定工作簿属性时间与 ZIP 条目顺序/时间戳/权限，
  便于后续内容寻址复用。

依赖：openpyxl（工程 Python 已可发现）。缺失时直接抛 ImportError，绝不降级为 CSV。
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import datetime
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

__all__ = [
    "export_operator_workbook",
    "WorkbookExportError",
    "ReadingDocumentError",
    "SCHEMA",
]

SCHEMA = "hf_decoder_reading.v1"
REQUIRED_PHASE_IDS = ("prefill", "decode")
SHEET_NAMES = {"prefill": "Prefill", "decode": "Decode"}

TENSOR_KINDS = ("activation", "weight", "cache", "mask", "index", "boundary")
SOURCE_KINDS = ("config", "default", "derived", "source", "runtime")
ANALYSIS_STATUSES = ("analyzed", "internal_unanalyzed")
STATUS_CN = {"analyzed": "已分析", "internal_unanalyzed": "内部未分析"}

# 稳定字节：固定属性时间与 ZIP 时间（ZIP 规范纪元 1980，取零时刻）。
_FIXED_EPOCH = datetime(1980, 1, 1, 0, 0, 0)
_FIXED_ZIP_DATE = (1980, 1, 1, 0, 0, 0)
_FIXED_ISO = "1980-01-01T00:00:00Z"
_FIXED_CREATOR = "structure_reader.export.operator_workbook"
# openpyxl 在 save 内部会把 properties.modified 覆写为保存时刻（秒级），
# 因此时间属性必须在稳定化阶段钉回固定值，否则字节随时间漂移、破坏内容寻址。
_CORE_TIME_RE = re.compile(r"(<dcterms:(created|modified)[^>]*>)(.*?)(</dcterms:\2>)", re.S)

_MAX_INLINE_REFS = 12   # 联查公式内最多展开的引用条数，超出只提示剩余项数
_FORMULA_LIMIT = 8000   # Excel 公式上限 8192，留余量后回退为纯文本


# ---------------------------------------------------------------- 异常


class WorkbookExportError(ValueError):
    """导出器错误基类（保持 ValueError 语义，调用方可按 ValueError 捕获）。"""


class ReadingDocumentError(WorkbookExportError):
    """输入阅读数据不符合冻结接口。"""


# ---------------------------------------------------------------- 校验工具


def _fail(path: str, message: str) -> None:
    raise ReadingDocumentError(f"{path}: {message}")


def _is_int(value) -> bool:
    return isinstance(value, int) and not isinstance(value, bool)


def _plain_str(value, path: str, label: str) -> str:
    if not isinstance(value, str):
        _fail(path, f"{label} 必须是字符串，实得 {type(value).__name__}")
    if ILLEGAL_CHARACTERS_RE.search(value):
        _fail(path, f"{label} 含 XML 非法控制字符，无法写入 xlsx")
    return value


def _require_mapping(value, path: str) -> dict:
    if not isinstance(value, dict):
        _fail(path, f"必须是对象(dict)，实得 {type(value).__name__}")
    return value


def _need(obj: dict, key: str, path: str):
    if key not in obj:
        _fail(path, f"缺少必需字段 {key!r}")
    return obj[key]


def _need_str(obj: dict, key: str, path: str) -> str:
    return _plain_str(_need(obj, key, path), f"{path}.{key}", key)


def _need_list(obj: dict, key: str, path: str) -> list:
    value = _need(obj, key, path)
    if not isinstance(value, list):
        _fail(f"{path}.{key}", f"必须是数组，实得 {type(value).__name__}")
    return value


def _str_list(obj: dict, key: str, path: str) -> list:
    items = _need_list(obj, key, path)
    return [_plain_str(item, f"{path}.{key}[{i}]", key) for i, item in enumerate(items)]


def _check_enum(value: str, allowed: tuple, path: str, label: str) -> str:
    if value not in allowed:
        _fail(path, f"{label} 必须是 {'/'.join(allowed)} 之一，实得 {value!r}")
    return value


def _json_value(value, path: str):
    """参数值允许 JSON 标量/数组/对象/null；其它 Python 类型一律拒绝。"""
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, list):
        return [_json_value(item, f"{path}[{i}]") for i, item in enumerate(value)]
    if isinstance(value, dict):
        checked = {}
        for key, item in value.items():
            name = _plain_str(str(key), f"{path}.{key}", "对象键")
            checked[name] = _json_value(item, f"{path}.{name}")
        return checked
    _fail(path, f"参数值必须是 JSON 标量/数组/对象或 null，实得 {type(value).__name__}")


def _canonical(value) -> str:
    """值的确定性字符串化：复合值用合法 JSON（对象键排序），保证字节稳定。"""
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return repr(value)
    if isinstance(value, str):
        return value
    if isinstance(value, (list, dict)):
        return json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def _value_kind(value) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "布尔"
    if isinstance(value, (int, float)):
        return "数值"
    if isinstance(value, str):
        return "文本"
    if isinstance(value, list):
        return "数组"
    if isinstance(value, dict):
        return "对象"
    return "未知"  # 校验后不可达


# ---------------------------------------------------------------- 文档解析


def _validate_source_refs(raw_list, path: str) -> list:
    refs = []
    for i, raw in enumerate(raw_list):
        item_path = f"{path}[{i}]"
        obj = _require_mapping(raw, item_path)
        refs.append(
            {
                "path": _need_str(obj, "path", item_path),
                "locator": _need_str(obj, "locator", item_path),
                "kind": _check_enum(
                    _need_str(obj, "kind", item_path), SOURCE_KINDS,
                    f"{item_path}.kind", "来源类型"),
                "note": _need_str(obj, "note", item_path),
            }
        )
    return refs


def _dedup(items: list) -> list:
    out = []
    for item in items:
        if item not in out:
            out.append(item)
    return out


def _validate_parameters(raw_list, path: str) -> list:
    params = []
    seen = {}
    for i, raw in enumerate(raw_list):
        item_path = f"{path}[{i}]"
        obj = _require_mapping(raw, item_path)
        pid = _need_str(obj, "id", item_path)
        if pid in seen:
            _fail(f"{item_path}.id", f"参数 id {pid!r} 在本阶段重复（先见于 {seen[pid]}）")
        seen[pid] = item_path
        expression = _need(obj, "expression", item_path)
        if expression is not None and not isinstance(expression, str):
            _fail(f"{item_path}.expression", "expression 必须是公式字符串或 null")
        params.append(
            {
                "id": pid,
                "symbol": _need_str(obj, "symbol", item_path),
                "meaning": _need_str(obj, "meaning", item_path),
                "value": _json_value(_need(obj, "value", item_path), f"{item_path}.value"),
                "expression": expression,
                "note": _need_str(obj, "note", item_path),
                "source_refs": _validate_source_refs(
                    _need_list(obj, "source_refs", item_path), f"{item_path}.source_refs"),
            }
        )
    return params


def _validate_shape(raw_shape, path: str) -> list:
    if not isinstance(raw_shape, list):
        _fail(path, f"shape 必须是按轴顺序的数组，实得 {type(raw_shape).__name__}")
    axes = []
    for i, axis in enumerate(raw_shape):
        axis_path = f"{path}[{i}]"
        if _is_int(axis):
            axes.append(axis)
        elif isinstance(axis, str):
            text = _plain_str(axis, axis_path, "shape 轴")
            if not text.strip():
                _fail(axis_path, "shape 轴表达式不能为空")
            axes.append(text)
        else:
            _fail(axis_path, "shape 轴必须是整数或符号表达式字符串，实得 "
                             f"{type(axis).__name__}")
    return axes


def _validate_tensors(raw_list, path: str) -> list:
    tensors = []
    seen = {}
    for i, raw in enumerate(raw_list):
        item_path = f"{path}[{i}]"
        obj = _require_mapping(raw, item_path)
        tid = _need_str(obj, "id", item_path)
        if tid in seen:
            _fail(f"{item_path}.id", f"张量 id {tid!r} 在本阶段重复（先见于 {seen[tid]}）")
        seen[tid] = item_path
        tensors.append(
            {
                "id": tid,
                "name": _need_str(obj, "name", item_path),
                "shape": _validate_shape(_need(obj, "shape", item_path), f"{item_path}.shape"),
                "parameter_refs": _str_list(obj, "parameter_refs", item_path),
                "kind": _check_enum(
                    _need_str(obj, "kind", item_path), TENSOR_KINDS,
                    f"{item_path}.kind", "张量类型"),
                "note": _need_str(obj, "note", item_path),
            }
        )
    return tensors


def _validate_operators(raw_list, path: str) -> list:
    operators = []
    seen = {}
    for i, raw in enumerate(raw_list):
        item_path = f"{path}[{i}]"
        obj = _require_mapping(raw, item_path)
        oid = _need_str(obj, "id", item_path)
        if oid in seen:
            _fail(f"{item_path}.id", f"算子 id {oid!r} 在本阶段重复（先见于 {seen[oid]}）")
        seen[oid] = item_path
        operators.append(
            {
                "id": oid,
                "module_path": _need_str(obj, "module_path", item_path),
                "name": _need_str(obj, "name", item_path),
                "op": _need_str(obj, "op", item_path),
                "inputs": _str_list(obj, "inputs", item_path),
                "outputs": _str_list(obj, "outputs", item_path),
                "weights": _str_list(obj, "weights", item_path),
                "parameter_refs": _str_list(obj, "parameter_refs", item_path),
                "source_refs": _validate_source_refs(
                    _need_list(obj, "source_refs", item_path), f"{item_path}.source_refs"),
                "analysis_status": _check_enum(
                    _need_str(obj, "analysis_status", item_path), ANALYSIS_STATUSES,
                    f"{item_path}.analysis_status", "分析状态"),
                "note": _need_str(obj, "note", item_path),
            }
        )
    return operators


def _resolve(refs: list, index: dict, kind: str, path: str) -> None:
    for ref in refs:
        if ref not in index:
            _fail(path, f"{kind}引用 {ref!r} 在本阶段内无法解析（不猜测、不补造）")


def _validate_phase(raw_phase, position: int) -> dict:
    path = f"phases[{position}]"
    obj = _require_mapping(raw_phase, path)
    pid = _need_str(obj, "id", f"{path}.id")
    expected = REQUIRED_PHASE_IDS[position]
    if pid != expected:
        _fail(f"{path}.id", f"第 {position + 1} 个阶段 id 必须是 {expected!r}，实得 {pid!r}")
    title = _need_str(obj, "title", f"{path}.title")
    params = _validate_parameters(_need_list(obj, "parameters", path), f"{path}.parameters")
    tensors = _validate_tensors(_need_list(obj, "tensors", path), f"{path}.tensors")
    operators = _validate_operators(_need_list(obj, "operators", path), f"{path}.operators")

    param_index = {p["id"]: p for p in params}
    tensor_index = {t["id"]: t for t in tensors}

    for tensor in tensors:
        _resolve(tensor["parameter_refs"], param_index, "参数",
                 f"张量 {tensor['id']}.parameter_refs")
    for op in operators:
        _resolve(op["inputs"], tensor_index, "输入张量", f"算子 {op['id']}.inputs")
        _resolve(op["outputs"], tensor_index, "输出张量", f"算子 {op['id']}.outputs")
        _resolve(op["weights"], tensor_index, "权重张量", f"算子 {op['id']}.weights")
        _resolve(op["parameter_refs"], param_index, "参数", f"算子 {op['id']}.parameter_refs")

    return {
        "id": pid,
        "title": title,
        "sheet_name": SHEET_NAMES[pid],
        "parameters": params,
        "tensors": tensors,
        "operators": operators,
        "cache_notes": _str_list(obj, "cache_notes", path),
        "coverage_notes": _str_list(obj, "coverage_notes", path),
        "param_index": param_index,
        "tensor_index": tensor_index,
    }


def _validate_document(reading_document) -> dict:
    if not isinstance(reading_document, dict):
        _fail("reading_document",
              f"必须是对象(dict)，实得 {type(reading_document).__name__}")
    doc = reading_document
    if _need_str(doc, "schema", "schema") != SCHEMA:
        _fail("schema", f"必须是 {SCHEMA!r}")

    model = _require_mapping(_need(doc, "model", "model"), "model")
    model_out = {
        key: _need_str(model, key, "model")
        for key in ("model_key", "display_name", "implementation", "source_revision", "config_path")
    }
    if model_out["implementation"] != "transformers":
        _fail("model.implementation",
              f"接口 v1 固定为 'transformers'，实得 {model_out['implementation']!r}")

    scope = _require_mapping(_need(doc, "scope", "scope"), "scope")
    layer_index = _need(scope, "layer_index", "scope")
    if not _is_int(layer_index) or layer_index < 0:
        _fail("scope.layer_index", "必须是非负整数")
    scope_out = {
        "layer_index": layer_index,
        "attention_type": _need_str(scope, "attention_type", "scope"),
        "moe_type": _need_str(scope, "moe_type", "scope"),
        "analysis_path": _need_str(scope, "analysis_path", "scope"),
        "notes": _str_list(scope, "notes", "scope"),
    }
    if scope_out["analysis_path"] != "eager":
        _fail("scope.analysis_path",
              f"接口 v1 固定为 'eager'，实得 {scope_out['analysis_path']!r}")

    raw_phases = _need_list(doc, "phases", "phases")
    if len(raw_phases) != 2:
        _fail("phases", "必须且只能是 prefill、decode 两个阶段，实得 "
                        f"{len(raw_phases)} 个（禁止第三个阶段或重复阶段）")
    ids = [_require_mapping(p, f"phases[{i}]").get("id") for i, p in enumerate(raw_phases)]
    if len(set(map(str, ids))) != 2:
        _fail("phases", f"阶段 id 重复：{ids!r}")

    return {
        "model": model_out,
        "scope": scope_out,
        "phases": [_validate_phase(p, i) for i, p in enumerate(raw_phases)],
    }


# ---------------------------------------------------------------- 样式与列定义

_THIN = Side(style="thin", color="FFB3BFC7")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_TOP = Alignment(vertical="top")
_CENTER = Alignment(vertical="center")

_FONT_TITLE = Font(bold=True, size=13, color="FFFFFFFF")
_FONT_SECTION = Font(bold=True, size=11, color="FFFFFFFF")
_FONT_HEAD = Font(bold=True, size=10, color="FF1F2933")
_FONT_LABEL = Font(bold=True, size=10, color="FF1F2933")
_FONT_BODY = Font(size=10, color="FF1F2933")
_FONT_BLACK = Font(bold=True, size=10, color="FF9C0006")
_FONT_LINK = Font(size=10, color="FF0563C1", underline="single")

_FILL_TITLE = PatternFill("solid", fgColor="FF1F4E79")
_FILL_PARAM_SECTION = PatternFill("solid", fgColor="FF215C36")
_FILL_OP_SECTION = PatternFill("solid", fgColor="FF7A3B0E")
_FILL_PARAM_HEAD = PatternFill("solid", fgColor="FFDDEBE0")
_FILL_OP_HEAD = PatternFill("solid", fgColor="FFF2E2D0")
_FILL_LABEL = PatternFill("solid", fgColor="FFEDF1F5")
_FILL_BLACK_CELL = PatternFill("solid", fgColor="FFFFC7CE")
_FILL_BLACK_ROW = PatternFill("solid", fgColor="FFFFF3F3")
_FILL_OK = PatternFill("solid", fgColor="FFE7F1E4")

PARAM_COLUMNS = (
    ("参数id", 18), ("符号", 14), ("含义", 30), ("值", 20), ("值类型", 10),
    ("表达式", 30), ("来源类型", 12), ("来源定位", 30), ("说明", 42),
)
OP_COLUMNS = (
    ("黑盒标记", 16), ("算子id", 18), ("模块路径", 30), ("名称", 22), ("算子类型", 16),
    ("输入 名称[shape]", 34), ("输出 名称[shape]", 30), ("权重 名称[shape]", 30),
    ("维度符号→本表配置值", 28), ("参数引用 本表id(可点击定位)", 24),
    ("参数联查 符号=值(公式指向本表参数区)", 34), ("源码/配置位置", 32),
    ("分析状态", 12), ("说明", 44),
)
N_COLS = len(OP_COLUMNS)
COLUMN_WIDTHS = tuple(
    max(
        PARAM_COLUMNS[i][1] if i < len(PARAM_COLUMNS) else 0,
        OP_COLUMNS[i][1],
    )
    for i in range(N_COLS)
)

DISCLAIMER = (
    "本表由 hf_reading_contract_v1 输入文档原样导出：结构与取值均来自该文档；"
    "导出器不推断模型语义、不填运行期假设值；文档中缺失/为空即呈现为空（标注未提供），未做补造。"
)


# ---------------------------------------------------------------- 单元格写入


def _text(ws, row: int, col: int, value, *, font=None, alignment=None, fill=None, border=True):
    """写文本单元格；把 openpyxl 会误判成公式/错误值的字符串锁回文本类型。"""
    cell = ws.cell(row=row, column=col)
    cell.value = "" if value is None else value
    if isinstance(value, str) or value is None:
        cell.data_type = "s"
    cell.font = font or _FONT_BODY
    cell.alignment = alignment or _WRAP_TOP
    if fill is not None:
        cell.fill = fill
    if border:
        cell.border = _BORDER
    return cell


def _number(ws, row: int, col: int, value: int | float):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.font = _FONT_BODY
    cell.alignment = _TOP
    cell.border = _BORDER
    return cell


def _formula(ws, row: int, col: int, formula: str, *, fill=None):
    cell = ws.cell(row=row, column=col)
    cell.value = formula          # 以 = 开头 → openpyxl 判定为公式（公式栏可见引用）
    cell.font = _FONT_BODY
    cell.alignment = _WRAP_TOP
    if fill is not None:
        cell.fill = fill
    cell.border = _BORDER
    return cell


def _banner(ws, row: int, text: str, fill, big: bool = False) -> None:
    for col in range(1, N_COLS + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = _FONT_TITLE if big else _FONT_SECTION
        cell.alignment = _CENTER
    ws.cell(row=row, column=1).value = text
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=N_COLS)
    ws.row_dimensions[row].height = 20 if big else 18


def _header_row(ws, row: int, columns, fill) -> None:
    for idx, (label, _w) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=idx)
        cell.value = label
        cell.font = _FONT_HEAD
        cell.alignment = _WRAP_TOP
        cell.fill = fill
        cell.border = _BORDER
    ws.row_dimensions[row].height = 32


# ---------------------------------------------------------------- 公式辅助


def _symbol_value_pieces(rows: list, ids: list | None = None) -> list:
    """把本表参数行的符号单元格与值单元格串成 $B$r&"="&$D$r。

    ids 与 rows 同序；条目非空时额外拼上该行的参数id单元格 $A$r（同名符号歧义时
    只用id能分辨哪行是哪个参数），不复制数值。
    """
    pieces = []
    for index, r in enumerate(rows):
        pid = None if ids is None else ids[index]
        if pid:
            pieces.append(f'$A${r}&":"&$B${r}&"="&$D${r}')
        else:
            pieces.append(f'$B${r}&"="&$D${r}')
    return pieces


def _xl_quote(text: str) -> str:
    return '"' + str(text).replace('"', '""') + '"'


def _lookup_formula(rows: list, ids: list | None = None,
                    suffix_text: str | None = None) -> str | None:
    """符号联查公式：值取自本表参数行单元格；suffix_text 作为字面量附注。

    无引用或超出 Excel 公式长度上限时返回 None（由调用方回退为纯文本）。
    """
    if not rows:
        return None
    shown = rows[:_MAX_INLINE_REFS]
    shown_ids = ids[:_MAX_INLINE_REFS] if ids is not None else None
    formula = "=" + ' & "; " & '.join(_symbol_value_pieces(shown, shown_ids))
    if len(rows) > len(shown):
        formula += f' & "; " & {_xl_quote(f"…另有{len(rows) - len(shown)}项，完整清单见「参数引用」列")}'
    if suffix_text:
        formula += f' & "; " & {_xl_quote(suffix_text)}'
    if len(formula) > _FORMULA_LIMIT:
        return None
    return formula


def _tensor_label(tensor: dict) -> str:
    axes = ", ".join(str(axis) for axis in tensor["shape"])
    return f"{tensor['name']}[{axes}]"


def _join_sources(refs: list) -> str:
    if not refs:
        return "（无来源）"
    parts = []
    for ref in refs:
        where = f"{ref['path']}#{ref['locator']}" if ref["path"] and ref["locator"] else (
            ref["path"] or ref["locator"] or "（未提供定位）")
        piece = f"[{ref['kind']}] {where}"
        if ref["note"]:
            piece += f"（{ref['note']}）"
        parts.append(piece)
    return " | ".join(parts)


# ---------------------------------------------------------------- 渲染


def _render_summary(ws, doc: dict, phase: dict) -> int:
    model, scope = doc["model"], doc["scope"]
    _banner(ws, 1,
            f"{model['display_name']} · {phase['title']}（decoder 第 {scope['layer_index']} 层阅读表）",
            _FILL_TITLE, big=True)
    facts = (
        ("阶段", f"{phase['id']} · {phase['title']}"),
        ("模型", f"{model['display_name']}（{model['model_key']}）"),
        ("实现", f"{model['implementation']} · source_revision {model['source_revision']}"),
        ("配置文件", model["config_path"]),
        ("阅读范围", f"layer_index={scope['layer_index']}；attention_type={scope['attention_type']}；"
                     f"moe_type={scope['moe_type']}；analysis_path={scope['analysis_path']}"),
        ("范围说明", " | ".join(scope["notes"]) or "（未提供）"),
        ("缓存说明", " | ".join(phase["cache_notes"]) or "（未提供）"),
        ("覆盖说明", " | ".join(phase["coverage_notes"]) or "（未提供）"),
        ("呈现约定", DISCLAIMER),
    )
    for offset, (label, value) in enumerate(facts):
        row = 2 + offset
        _text(ws, row, 1, label, font=_FONT_LABEL, alignment=_TOP, fill=_FILL_LABEL, border=False)
        _text(ws, row, 2, value, alignment=_TOP, border=False)
    return 2 + len(facts) + 1  # 空一行分隔下一区


def _render_parameters(ws, start_row: int, phase: dict):
    """参数符号区。返回 (下一空闲行, 参数id→行号, 符号→本表同名行列表, 行号→参数id)。

    接口只保证参数 id 阶段内唯一，symbol 可重复；因此符号必须映射到**全部**候选行，
    不能“取第一个”默认绑定（同名错绑会静默给出错误配置值）。
    """
    _banner(ws, start_row,
            f"参数符号区（{len(phase['parameters'])} 项；算子与张量的 parameter_refs 引用本区，"
            f"值/含义/来源均可在此核对）", _FILL_PARAM_SECTION)
    head_row = start_row + 1
    _header_row(ws, head_row, PARAM_COLUMNS, _FILL_PARAM_HEAD)

    row_of_id: dict = {}
    rows_by_symbol: dict = {}
    id_of_row: dict = {}
    row = head_row + 1
    for param in phase["parameters"]:
        row_of_id[param["id"]] = row
        id_of_row[row] = param["id"]
        rows_by_symbol.setdefault(param["symbol"], []).append(row)
        _text(ws, row, 1, param["id"])
        _text(ws, row, 2, param["symbol"], font=_FONT_LABEL)
        _text(ws, row, 3, param["meaning"])
        value = param["value"]
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            _number(ws, row, 4, value)
        else:
            _text(ws, row, 4, _canonical(value))
        _text(ws, row, 5, _value_kind(value))
        _text(ws, row, 6, param["expression"] or "（无表达式）")
        _text(ws, row, 7, ", ".join(sorted({r["kind"] for r in param["source_refs"]}))
              or "（未提供）")
        _text(ws, row, 8, _join_sources(param["source_refs"]))
        _text(ws, row, 9, param["note"])
        row += 1
    if not phase["parameters"]:
        _text(ws, row, 1, "（本阶段文档未提供参数）", font=_FONT_LABEL)
        row += 1
    return row + 1, row_of_id, rows_by_symbol, id_of_row


def _is_symbolic_axis(axis) -> bool:
    """只有真正的符号/表达式轴才参与参数区联查；纯数字字面量（含字符串形式）是具体尺度。"""
    return isinstance(axis, str) and not re.fullmatch(r"-?\d+", axis.strip())


def _axis_symbols(operator: dict, tensor_index: dict) -> list:
    symbols: list = []
    for tid in (*operator["inputs"], *operator["outputs"], *operator["weights"]):
        for axis in tensor_index[tid]["shape"]:
            if _is_symbolic_axis(axis) and axis not in symbols:
                symbols.append(axis)
    return symbols


def _operator_param_ids(operator: dict, tensor_index: dict) -> list:
    """算子显式引用 + 所用张量显式引用（去重保序，只用文档声明的引用）。"""
    involved = [tensor_index[tid] for tid in
                (*operator["inputs"], *operator["outputs"], *operator["weights"])]
    return _dedup([*operator["parameter_refs"],
                   *(pid for tensor in involved for pid in tensor["parameter_refs"])])


def _bind_axis_symbols(symbols: list, rows_by_symbol: dict, id_of_row: dict,
                       explicit_ids: set):
    """用显式 parameter_refs 定位轴符号对应的本表参数行，同名字母不能默认取第一个。

    返回 (已绑定行, 歧义符号及其同名候选, 本表无同名参数的符号)：
    - 同名候选只有一个 → 直接绑定（原有行为）。
    - 多个同名候选 → 只在当前行（算子及其输入/输出/权重张量）的显式引用能唯一分辨时绑定。
    - 显式引用分辨不出（零个或多个候选被引用）→ 记为歧义，不猜值、不静静选错。
    """
    bound: list = []
    ambiguous: list = []
    missed: list = []
    for symbol in symbols:
        candidates = rows_by_symbol.get(symbol) or []
        if not candidates:
            missed.append(symbol)
            continue
        if len(candidates) == 1:
            bound.append(candidates[0])
            continue
        hits = [r for r in candidates if id_of_row[r] in explicit_ids]
        if len(hits) == 1:
            bound.append(hits[0])
        else:
            ambiguous.append((symbol, [(id_of_row[r], r) for r in (hits or candidates)]))
    return bound, ambiguous, missed


def _unbound_note(ambiguous: list, missed: list) -> str | None:
    """歧义/未命中的中文陈达：列出相关参数id与本表行号，供人工定位。"""
    parts = []
    if ambiguous:
        bodies = [f"{symbol}（本表同名候选："
                  + "、".join(f"{pid}（第{r}行）" for pid, r in pairs) + "）"
                  for symbol, pairs in ambiguous]
        parts.append("符号存在歧义，未取配置值：" + "；".join(bodies))
    if missed:
        parts.append("未命中本表参数区同名符号：" + ", ".join(missed))
    return "；".join(parts) if parts else None


def _render_operators(ws, start_row: int, phase: dict, row_of_id: dict,
                      rows_by_symbol: dict, id_of_row: dict) -> None:
    sheet_name = phase["sheet_name"]
    tensor_index = phase["tensor_index"]
    _banner(ws, start_row,
            f"算子区（{len(phase['operators'])} 行，按实际阅读顺序；"
            f"◆ 黑盒 = 内部未分析，参数/维度联查列指向本表参数区）", _FILL_OP_SECTION)
    head_row = start_row + 1
    _header_row(ws, head_row, OP_COLUMNS, _FILL_OP_HEAD)

    row = head_row + 1
    for op in phase["operators"]:
        blackbox = op["analysis_status"] == "internal_unanalyzed"
        row_fill = _FILL_BLACK_ROW if blackbox else None
        status_font = _FONT_BLACK if blackbox else _FONT_BODY

        def put(col: int, value, font=None):
            return _text(ws, row, col, value,
                         font=font or (status_font if blackbox and col in (1, 13) else _FONT_BODY),
                         fill=row_fill)

        put(1, "◆ 黑盒（内部未分析）" if blackbox else "已展开")
        put(2, op["id"])
        put(3, op["module_path"])
        put(4, op["name"])
        put(5, op["op"])

        def group(ids: list) -> str:
            return ", ".join(_tensor_label(tensor_index[t]) for t in ids) if ids else "（无）"

        put(6, group(op["inputs"]))
        put(7, group(op["outputs"]))
        put(8, group(op["weights"]))

        # 显式引用集是本表绑定权威：先算出来，再拿它分辨轴符号。
        ref_ids = _operator_param_ids(op, tensor_index)
        ref_rows = [row_of_id[pid] for pid in ref_ids]

        symbols = _axis_symbols(op, tensor_index)
        bound_rows, ambiguous, missed = _bind_axis_symbols(
            symbols, rows_by_symbol, id_of_row, set(ref_ids))
        note = _unbound_note(ambiguous, missed)
        formula = _lookup_formula(bound_rows, suffix_text=note)
        if formula is None and bound_rows:
            formula = _lookup_formula(bound_rows)  # 附注导致超长时至少保留已命中的联查
        if formula:
            _formula(ws, row, 9, formula, fill=row_fill)
        else:
            put(9, note if note else "（无符号轴）")

        link = put(10, ", ".join(ref_ids) if ref_ids else "（无参数引用）")
        if ref_ids:
            target = row_of_id[ref_ids[0]]
            link.hyperlink = Hyperlink(
                ref=link.coordinate,
                location=f"'{sheet_name}'!$A${target}",
                tooltip=f"定位本表参数区 {ref_ids[0]}（第 {target} 行）",
                display=link.value,
            )
            link.font = _FONT_LINK
        # 引用的参数中存在同名符号时，联查附id（$A$ 同行），否则保持原有简洁形式。
        ref_symbols = [phase["param_index"][pid]["symbol"] for pid in ref_ids]
        dup_symbols = {s for s in ref_symbols if ref_symbols.count(s) > 1}
        k_ids = [pid if s in dup_symbols else None
                 for pid, s in zip(ref_ids, ref_symbols)]
        lookup = _lookup_formula(ref_rows, k_ids)
        if lookup is None and ref_rows:
            # 公式过长：只展开前 4 项，其余降为数量提示（不复制无法联查的值）。
            lookup = _lookup_formula(ref_rows[:4], k_ids[:4],
                                     f"…另有{len(ref_rows) - 4}项，完整清单见「参数引用」列")
        if lookup:
            _formula(ws, row, 11, lookup, fill=row_fill)
        else:
            put(11, "（无参数引用）" if not ref_rows else "（引用无法展开，见「参数引用」列）")

        put(12, _join_sources(op["source_refs"]))
        status = put(13, STATUS_CN[op["analysis_status"]])
        status.fill = _FILL_BLACK_CELL if blackbox else _FILL_OK
        status.font = _FONT_BLACK if blackbox else _FONT_BODY
        put(14, op["note"])
        row += 1

    if not phase["operators"]:
        _text(ws, row, 1, "（本阶段文档未提供算子）", font=_FONT_LABEL)
        row += 1
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(N_COLS)}{row - 1}"


def _build_sheet(ws, doc: dict, phase: dict) -> None:
    for idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    next_row = _render_summary(ws, doc, phase)
    next_row, row_of_id, rows_by_symbol, id_of_row = _render_parameters(ws, next_row, phase)
    _render_operators(ws, next_row, phase, row_of_id, rows_by_symbol, id_of_row)
    ws.freeze_panes = "B2"          # 首列 id 与标题横幅常驻
    ws.sheet_properties.tabColor = "FF1F4E79" if phase["id"] == "prefill" else "FF215C36"


# ---------------------------------------------------------------- ZIP 稳定化


def _stabilize_core_props(payload: bytes) -> bytes:
    """把 core.xml 的 created/modified 钉为固定时刻（不静默放过开放形式）。"""
    text = payload.decode("utf-8")
    patched, count = _CORE_TIME_RE.subn(lambda m: m.group(1) + _FIXED_ISO + m.group(4), text)
    if count != 2:
        raise WorkbookExportError(f"无法固定工作簿时间属性（匹配 {count} 处，预期 2 处）")
    return patched.encode("utf-8")


def _normalize_zip(payload: bytes) -> bytes:
    """固定条目顺序、时间戳、权限、创建系统与属性时间，保证相同输入字节一致。"""
    source = zipfile.ZipFile(io.BytesIO(payload), "r")
    try:
        names = sorted(source.namelist(), key=lambda n: (n != "[Content_Types].xml", n))
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
            for name in names:
                data = source.read(name)
                if name == "docProps/core.xml":
                    data = _stabilize_core_props(data)
                info = zipfile.ZipInfo(name, date_time=_FIXED_ZIP_DATE)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = 0o600 << 16
                info.internal_attr = 0
                info.comment = b""
                target.writestr(info, data)
        return buffer.getvalue()
    finally:
        source.close()


def export_operator_workbook(reading_document: dict) -> bytes:
    """按 hf_reading_contract_v1 把阅读数据导出为 Prefill/Decode 两表 xlsx 字节。

    只读输入（不修改调用方对象）、只返回字节（不写文件、不读磁盘、不导入 HF 提取器）。
    输入不合法时抛 ReadingDocumentError（ValueError 子类）。
    """
    doc = _validate_document(reading_document)

    wb = Workbook()
    wb.iso_dates = True
    props = wb.properties
    props.creator = _FIXED_CREATOR
    props.lastModifiedBy = _FIXED_CREATOR
    props.title = f"{doc['model']['display_name']} decoder 阅读表"
    props.description = ("schema hf_decoder_reading.v1；仅 Prefill/Decode 两表；"
                         "由 structure_reader.export.operator_workbook 导出")
    props.subject = f"{doc['model']['model_key']} layer {doc['scope']['layer_index']}"
    props.created = _FIXED_EPOCH
    props.modified = _FIXED_EPOCH

    sheets = []
    first = True
    for phase in doc["phases"]:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = phase["sheet_name"]
        _build_sheet(ws, doc, phase)
        sheets.append(ws.title)
    wb.active = 0
    if sheets != ["Prefill", "Decode"]:
        raise WorkbookExportError(f"内部错误：表顺序异常 {sheets!r}")

    buffer = io.BytesIO()
    wb.save(buffer)
    return _normalize_zip(buffer.getvalue())
