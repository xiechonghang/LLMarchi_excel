"""两表 Excel 导出器自包含测试（pi 侧，hf_reading_contract_v1）。

测试数据：tests/hf/fixtures/workbook_synthetic_document.json —— 明确为接口合成数据，
只用于验证导出契约（两表、字段、引用联查、黑盒标注、拒绝非法输入、字节稳定）。
它不是真实 DeepSeek 产物、不是 golden、不代表任何模型事实；真实模型语义由 Claude
依 HF 源码构建并由 opencode 独立验证。

本文件不读取 models/、outputs/、.work/，不依赖 Claude/opencode 的模块或产物，
不访问网络；只用工程 Python 已存在的 openpyxl 与标准库 unittest。
运行：python3 -m unittest discover -s structure_reader/tests/hf -t structure_reader/tests
"""

from __future__ import annotations

import copy
import io
import json
import re
import sys
import time
import unittest
import zipfile
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
SRC_DIR = HERE.parent.parent / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from structure_reader.export.operator_workbook import (  # noqa: E402
    ReadingDocumentError,
    WorkbookExportError,
    export_operator_workbook,
)

from openpyxl import load_workbook  # noqa: E402

FIXTURE = HERE / "fixtures" / "workbook_synthetic_document.json"

PARAM_HEADER = "参数id"
OP_HEADER = "黑盒标记"
COL = {  # 算子区列位（1 基）
    "marker": 1, "op_id": 2, "module": 3, "name": 4, "op": 5,
    "inputs": 6, "outputs": 7, "weights": 8, "axis_lookup": 9,
    "param_ids": 10, "param_lookup": 11, "source": 12, "status": 13, "note": 14,
}
PARAM_COL = {"id": 1, "symbol": 2, "meaning": 3, "value": 4, "value_kind": 5,
             "expression": 6, "source_kind": 7, "source_locator": 8, "note": 9}
N_COLS = 14


def load_doc() -> dict:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def sheet_of(payload: bytes, name: str):
    wb = load_workbook(io.BytesIO(payload))
    return wb, wb[name]


def _col_a(ws, row: int):
    value = ws.cell(row=row, column=1).value
    return value if isinstance(value, str) else ""


def param_data_rows(ws) -> list:
    """参数区数据行：参数表头之后，直到区隔空行/算子区横幅。"""
    header = next(r for r in range(1, ws.max_row + 1) if _col_a(ws, r) == PARAM_HEADER)
    rows = []
    for r in range(header + 1, ws.max_row + 1):
        a = _col_a(ws, r)
        if a == "" or a.startswith("算子区"):
            break
        rows.append(r)
    return rows


def op_data_rows(ws) -> list:
    """算子区数据行：算子表头之后，直到空行/末尾。"""
    header = next(r for r in range(1, ws.max_row + 1) if _col_a(ws, r) == OP_HEADER)
    rows = []
    for r in range(header + 1, ws.max_row + 1):
        if _col_a(ws, r) == "":
            break
        rows.append(r)
    return rows


def cell(ws, row: int, key: str) -> str:
    return str(ws.cell(row=row, column=COL[key]).value or "")


def param_cell(ws, row: int, key: str) -> str:
    return str(ws.cell(row=row, column=PARAM_COL[key]).value or "")


def formula_targets(formula: str) -> list:
    """公式里出现的本表单元格引用：[(列, 行), ...]。"""
    return re.findall(r"\$([A-Z]+)\$(\d+)", formula or "")


class WorkbookExportTestBase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.doc = load_doc()
        cls.payload = export_operator_workbook(cls.doc)

    def mutate(self, editor) -> dict:
        """在合成文档副本上做一处改动，返回改后的文档（不改类级缓存）。"""
        doc = copy.deepcopy(self.doc)
        editor(doc)
        return doc


class TestWorkbookShape(WorkbookExportTestBase):
    def test_returns_valid_xlsx_bytes(self):
        self.assertIsInstance(self.payload, bytes)
        self.assertGreater(len(self.payload), 4000)
        with zipfile.ZipFile(io.BytesIO(self.payload)) as zf:
            self.assertIsNone(zf.testzip())
            self.assertIn("xl/workbook.xml", zf.namelist())

    def test_exactly_two_sheets_prefill_decode(self):
        wb, _ = sheet_of(self.payload, "Prefill")
        self.assertEqual(wb.sheetnames, ["Prefill", "Decode"])
        self.assertEqual(len(wb.worksheets), 2)
        for ws in wb.worksheets:  # 不允许隐藏 metadata sheet
            self.assertEqual(ws.sheet_state, "visible")

    def test_each_sheet_is_self_contained(self):
        for name in ("Prefill", "Decode"):
            _, ws = sheet_of(self.payload, name)
            self.assertIn("合成接口探针模型", cell(ws, 1, "marker"))       # 第 1 行横幅
            self.assertTrue(param_data_rows(ws), name)
            self.assertTrue(op_data_rows(ws), name)
            self.assertEqual(ws.freeze_panes, "B2")
            self.assertTrue(ws.auto_filter.ref.startswith("A"))
            self.assertGreaterEqual(ws.max_column, N_COLS)

    def test_no_repetition_count_column_or_per_expert_rows(self):
        for name in ("Prefill", "Decode"):
            _, ws = sheet_of(self.payload, name)
            texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
            joined = " ".join(texts).lower()
            for banned in ("repeat_count", "repetition_count", "重复次数"):
                self.assertNotIn(banned, joined)
            # 逐专家复制：算子行数必须等于文档算子数（K 个专家不展开成 K 行）
            self.assertEqual(len(op_data_rows(ws)), len(_phase(name)["operators"]))


class TestParameterArea(WorkbookExportTestBase):
    def test_parameter_area_has_required_fields(self):
        for name in ("Prefill", "Decode"):
            _, ws = sheet_of(self.payload, name)
            phase = _phase(name)
            rows = param_data_rows(ws)
            self.assertEqual(len(rows), len(phase["parameters"]))
            ids = [param_cell(ws, r, "id") for r in rows]
            self.assertEqual(ids, [p["id"] for p in phase["parameters"]])
            for r, param in zip(rows, phase["parameters"]):
                self.assertEqual(param_cell(ws, r, "symbol"), param["symbol"])
                self.assertEqual(param_cell(ws, r, "meaning"), param["meaning"])
                self.assertTrue(param_cell(ws, r, "expression"))
                self.assertTrue(param_cell(ws, r, "source_kind"))
                self.assertTrue(param_cell(ws, r, "source_locator"))
                self.assertEqual(param_cell(ws, r, "note"), param["note"])

    def test_values_are_not_fabricated(self):
        _, ws = sheet_of(self.payload, "Prefill")
        rows = param_data_rows(ws)
        by_id = {param_cell(ws, r, "id"): r for r in rows}
        # 运行期值：文档为 null，表格必须呈现 null，不能编出 B/T 的固定数字
        self.assertEqual(param_cell(ws, by_id["p_batch"], "value"), "null")
        self.assertEqual(param_cell(ws, by_id["p_seq"], "value"), "null")
        self.assertEqual(param_cell(ws, by_id["p_batch"], "value_kind"), "null")
        # 标量数值保持为数值单元格（可比较），复合值确定性字符串化为合法 JSON
        self.assertEqual(ws.cell(row=by_id["p_layers"], column=PARAM_COL["value"]).value, 43)
        self.assertEqual(param_cell(ws, by_id["p_rope"], "value"),
                         '{"beta_fast":32,"beta_slow":1,"mscale":1.0}')
        self.assertEqual(param_cell(ws, by_id["p_router_bias"], "value"), "[0.0,0.25,1.5]")

    def test_text_starting_with_equals_stays_text(self):
        _, ws = sheet_of(self.payload, "Prefill")
        by_id = {param_cell(ws, r, "id"): r for r in param_data_rows(ws)}
        cell_obj = ws.cell(row=by_id["p_hidden"], column=PARAM_COL["expression"])
        self.assertEqual(cell_obj.value, "= 2 * config.intermediate_size_hint")
        self.assertEqual(cell_obj.data_type, "s")  # 不是公式


class TestOperatorRows(WorkbookExportTestBase):
    def test_operator_row_fields_in_reading_order(self):
        _, ws = sheet_of(self.payload, "Prefill")
        rows = op_data_rows(ws)
        phase = _phase("prefill")
        self.assertEqual([cell(ws, r, "op_id") for r in rows], [o["id"] for o in phase["operators"]])
        for r, op in zip(rows, phase["operators"]):
            self.assertEqual(cell(ws, r, "module"), op["module_path"])
            self.assertEqual(cell(ws, r, "name"), op["name"])
            self.assertEqual(cell(ws, r, "op"), op["op"])
            self.assertTrue(cell(ws, r, "source"))
            self.assertEqual(cell(ws, r, "note"), op["note"])

    def test_shapes_keep_symbols_and_weight_names(self):
        _, ws = sheet_of(self.payload, "Prefill")
        rows = {cell(ws, r, "op_id"): r for r in op_data_rows(ws)}
        router = cell(ws, rows["op_router"], "weights")
        self.assertIn("gate_weight[E, H]", router)          # 符号保留
        self.assertIn("hidden_states[B, T, H]", cell(ws, rows["op_router"], "inputs"))
        self.assertNotIn("4096", cell(ws, rows["op_router"], "inputs"))  # 不代入配置值
        self.assertEqual(cell(ws, rows["op_embed_lookup"], "weights"), "（无）")
        # 标量张量 shape 为空的呈现
        self.assertIn("routed_gain[]", cell(ws, rows["op_routed_expert"], "inputs"))

    def test_tensor_parameter_refs_are_aggregated(self):
        _, ws = sheet_of(self.payload, "Prefill")
        rows = {cell(ws, r, "op_id"): r for r in op_data_rows(ws)}
        refs = cell(ws, rows["op_router"], "param_ids").split(", ")
        # 算子自身 p_topk + 张量 hidden_states 的 p_batch/p_seq/p_hidden + gate_weight 的 p_hidden
        self.assertIn("p_topk", refs)
        self.assertIn("p_hidden", refs)
        self.assertEqual(refs.count("p_hidden"), 1)
        self.assertEqual(cell(ws, rows["op_embed_lookup"], "param_ids").split(", ").count("p_seq"), 1)

    def test_param_refs_link_to_this_sheet_parameter_area(self):
        for name in ("Prefill", "Decode"):
            _, ws = sheet_of(self.payload, name)
            prow = param_data_rows(ws)
            param_ids = {param_cell(ws, r, "id"): r for r in prow}
            for r in op_data_rows(ws):
                refs = ([] if cell(ws, r, "param_ids").startswith("（")
                        else cell(ws, r, "param_ids").split(", "))
                link = ws.cell(row=r, column=COL["param_ids"]).hyperlink
                if not refs:
                    self.assertIsNone(link, (name, r))
                    continue
                # 点击定位：链接必须落在本表参数区且指向第一个引用 id
                self.assertEqual(link.location, f"'{name}'!$A${param_ids[refs[0]]}")
                # 联查公式：只能是本表参数区单元格（B=符号、D=值；同名歧义时附 A=参数id）
                formula = cell(ws, r, "param_lookup")
                self.assertTrue(formula.startswith("="), (name, r))
                for col_letter, row_ref in formula_targets(formula):
                    self.assertIn(col_letter, ("A", "B", "D"))
                    self.assertIn(int(row_ref), prow, (name, r, row_ref))


class TestLookupFormulas(WorkbookExportTestBase):
    def test_axis_symbol_lookup_resolves_values(self):
        _, ws = sheet_of(self.payload, "Prefill")
        rows = {cell(ws, r, "op_id"): r for r in op_data_rows(ws)}
        prow = {param_cell(ws, r, "symbol"): r for r in param_data_rows(ws)}
        formula = cell(ws, rows["op_embed_lookup"], "axis_lookup")
        self.assertTrue(formula.startswith("="))
        # T 命中本表参数区 p_seq 行：公式引用该行的 B/D 单元格
        self.assertIn(("B", str(prow["T"])), formula_targets(formula))
        # E 在本表参数区无同名符号：必须写明未命中，而不是猜值
        shared = cell(ws, rows["op_shared_expert"], "axis_lookup")
        self.assertIn("未命中本表参数区同名符号：N", shared)

    def test_numeric_axis_string_is_not_treated_as_symbol(self):
        _, ws = sheet_of(self.payload, "Decode")
        rows = {cell(ws, r, "op_id"): r for r in op_data_rows(ws)}
        self.assertNotIn("未命中本表参数区同名符号：1",
                         cell(ws, rows["op_moe_step"], "axis_lookup"))

    def test_weight_lookup_value_and_symbol_both_readable(self):
        """shape 符号可读到对应配置值：联查公式取自参数行单元格。"""
        _, ws = sheet_of(self.payload, "Prefill")
        rows = {cell(ws, r, "op_id"): r for r in op_data_rows(ws)}
        prow = {param_cell(ws, r, "symbol"): r for r in param_data_rows(ws)}
        targets = set(formula_targets(cell(ws, rows["op_router"], "param_lookup")))
        self.assertIn(("D", str(prow["K"])), targets)   # K 的值取自参数区 D 列
        self.assertIn(("D", str(prow["H"])), targets)   # 权重 gate_weight[E, H] 的 H


class TestBlackboxMarking(WorkbookExportTestBase):
    def test_internal_unanalyzed_marked_in_chinese(self):
        phase = _phase("prefill")
        blackbox = {o["id"] for o in phase["operators"]
                    if o["analysis_status"] == "internal_unanalyzed"}
        self.assertTrue(blackbox)
        _, ws = sheet_of(self.payload, "Prefill")
        rows = {cell(ws, r, "op_id"): r for r in op_data_rows(ws)}
        for oid in blackbox:
            r = rows[oid]
            self.assertEqual(cell(ws, r, "status"), "内部未分析")
            self.assertIn("黑盒", cell(ws, r, "marker"))
            status = ws.cell(row=r, column=COL["status"])
            self.assertTrue(status.font.b)
            self.assertEqual(status.fill.fgColor.rgb, "FFFFC7CE")
            self.assertEqual(status.alignment.wrap_text, True)
        for oid in set(rows) - blackbox:
            self.assertEqual(cell(ws, rows[oid], "status"), "已分析")

    def test_raw_enum_never_leaks_into_sheets(self):
        for name in ("Prefill", "Decode"):
            _, ws = sheet_of(self.payload, name)
            texts = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
            self.assertNotIn("internal_unanalyzed", texts)


class TestSharedTensorReferences(WorkbookExportTestBase):
    def test_shared_tensor_referenced_twice_without_duplicate_rows(self):
        phase = _phase("prefill")
        users = [o["id"] for o in phase["operators"] if "t_shared_weight" in o["weights"]]
        self.assertEqual(len(users), 2, "合成数据应包含共享权重被两个算子引用")
        _, ws = sheet_of(self.payload, "Prefill")
        rows = {cell(ws, r, "op_id"): r for r in op_data_rows(ws)}
        for oid in users:
            self.assertIn("shared_expert_weight[N, H]", cell(ws, rows[oid], "weights"))
        # 共享引用不是重复定义：参数区行数与算子区行数不因复用而增加
        self.assertEqual(len(param_data_rows(ws)), len(phase["parameters"]))
        self.assertEqual(len(op_data_rows(ws)), len(phase["operators"]))

    def test_same_id_in_different_phases_is_phase_scoped(self):
        prefill_ids = {p["id"] for p in _phase("prefill")["parameters"]}
        decode_ids = {p["id"] for p in _phase("decode")["parameters"]}
        self.assertTrue(prefill_ids & decode_ids)
        _, wsf = sheet_of(self.payload, "Prefill")
        _, wsd = sheet_of(self.payload, "Decode")
        self.assertNotEqual(len(param_data_rows(wsf)), len(param_data_rows(wsd)))
        # 两表各自独立成文：Decode 的 T=1 不与 Prefill 的 T=null 混用
        def value_of(ws, pid):
            rows = param_data_rows(ws)
            r = next(x for x in rows if param_cell(ws, x, "id") == pid)
            return param_cell(ws, r, "symbol"), param_cell(ws, r, "value")
        self.assertEqual(value_of(wsf, "p_seq"), ("T", "null"))
        self.assertEqual(value_of(wsd, "p_step"), ("T", "1"))


class TestDeterministicBytes(WorkbookExportTestBase):
    def test_same_input_same_bytes(self):
        first = export_operator_workbook(copy.deepcopy(self.doc))
        second = export_operator_workbook(copy.deepcopy(self.doc))
        self.assertEqual(first, second)

    def test_bytes_do_not_drift_with_wall_clock(self):
        """回归：openpyxl 会把 dcterms:modified 写成保存时刻，必须被钉回固定值。"""
        first = export_operator_workbook(copy.deepcopy(self.doc))
        time.sleep(1.05)
        second = export_operator_workbook(copy.deepcopy(self.doc))
        self.assertEqual(first, second)
        self.assertEqual(first, self.payload)

    def test_time_pinning_fails_loudly_if_openpyxl_layout_changes(self):
        """不能静默放过时间属性：钉不住就报错，而不是默默产出漂移字节。"""
        import structure_reader.export.operator_workbook as mod

        with self.assertRaises(WorkbookExportError):
            mod._stabilize_core_props(b"<cp:coreProperties><dc:creator>x</dc:creator></cp:coreProperties>")

    def test_key_order_of_input_does_not_change_bytes(self):
        def reversed_keys(obj):
            if isinstance(obj, dict):
                return {k: reversed_keys(v) for k, v in reversed(list(obj.items()))}
            if isinstance(obj, list):
                return [reversed_keys(v) for v in obj]
            return obj

        payload = export_operator_workbook(reversed_keys(copy.deepcopy(self.doc)))
        self.assertEqual(payload, self.payload)

    def test_volatile_zip_metadata_pinned(self):
        with zipfile.ZipFile(io.BytesIO(self.payload)) as zf:
            infos = zf.infolist()
            self.assertEqual({i.date_time for i in infos}, {(1980, 1, 1, 0, 0, 0)})
            self.assertEqual({i.create_system for i in infos}, {0})
            self.assertEqual({i.external_attr for i in infos}, {0o600 << 16})
            self.assertEqual([i.filename for i in infos],
                             sorted([i.filename for i in infos],
                                    key=lambda n: (n != "[Content_Types].xml", n)))
            core = zf.read("docProps/core.xml").decode("utf-8")
            stamps = re.findall(r"<dcterms:(?:created|modified)[^>]*>([^<]+)</dcterms:", core)
            self.assertEqual(stamps, ["1980-01-01T00:00:00Z", "1980-01-01T00:00:00Z"], core)
            self.assertNotIn(datetime.now().strftime("%Y-%m-%dT"), core)

    def test_formulas_survive_the_file(self):
        """联查靠真公式：落盘后 XML 里必须是 <f> 元素，不是展平后的常量。"""
        with zipfile.ZipFile(io.BytesIO(self.payload)) as zf:
            for part in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
                xml = zf.read(part).decode("utf-8")
                self.assertIn("<f>", xml, part)
                self.assertIn("<hyperlink ", xml, part)

    def test_output_is_not_csv_or_markdown(self):
        self.assertTrue(self.payload.startswith(b"PK\x03\x04"))
        self.assertNotIn(b",,", self.payload[:2000])


class TestInputIsNotConsumedOrRewritten(unittest.TestCase):
    def test_document_object_is_not_modified(self):
        doc = load_doc()
        snapshot = json.dumps(doc, sort_keys=True, ensure_ascii=False)
        export_operator_workbook(doc)
        self.assertEqual(json.dumps(doc, sort_keys=True, ensure_ascii=False), snapshot)

    def test_exporter_only_touches_bytes_and_openpyxl(self):
        """静态不变量：不读磁盘/不起进程/不导入 HF 提取器或 torch。"""
        import ast

        import structure_reader.export.operator_workbook as mod

        tree = ast.parse(Path(mod.__file__).read_text(encoding="utf-8"))
        roots = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                roots.update(a.name.split(".")[0] for a in node.names)
            elif isinstance(node, ast.ImportFrom):
                if node.module:
                    roots.add(node.module.split(".")[0])
        self.assertLessEqual(roots, {"__future__", "io", "json", "re", "zipfile",
                                     "datetime", "openpyxl"}, sorted(roots))
        calls = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Call):
                if isinstance(node.func, ast.Name):
                    calls.add(node.func.id)
                elif isinstance(node.func, ast.Attribute):
                    calls.add(node.func.attr)
        for banned in ("open", "system", "popen", "run", "check_output", "urlopen",
                       "download_file", "exec", "eval"):
            self.assertNotIn(banned, calls, banned)


class TestRejectedDocuments(unittest.TestCase):
    def setUp(self):
        self.doc = load_doc()

    def expect_reject(self, editor, needle: str):
        doc = copy.deepcopy(self.doc)
        editor(doc)
        with self.assertRaises(ValueError) as ctx:
            export_operator_workbook(doc)
        self.assertIsInstance(ctx.exception, ReadingDocumentError)
        self.assertIn(needle, str(ctx.exception))

    def test_non_dict_input(self):
        for bad in (None, [], "x", 7):
            with self.assertRaises(ValueError):
                export_operator_workbook(bad)

    def test_bad_schema(self):
        self.expect_reject(lambda d: d.update(schema="hf_decoder_reading.v2"), "schema")

    def test_third_phase_rejected(self):
        self.expect_reject(lambda d: d["phases"].append(copy.deepcopy(d["phases"][0])), "两个阶段")

    def test_duplicate_phase_id_rejected(self):
        def edit(doc):
            doc["phases"][1]["id"] = "prefill"

        self.expect_reject(edit, "阶段")

    def test_phase_order_rejected(self):
        def edit(doc):
            doc["phases"] = list(reversed(doc["phases"]))

        self.expect_reject(edit, "prefill")

    def test_missing_required_field_rejected(self):
        def edit(doc):
            del doc["phases"][0]["operators"][0]["op"]

        self.expect_reject(edit, "op")

    def test_duplicate_ids_rejected(self):
        for path, key, needle in (
            ("parameters", "id", "参数 id"),
            ("tensors", "id", "张量 id"),
            ("operators", "id", "算子 id"),
        ):
            def edit(doc, path=path, key=key):
                doc["phases"][0][path][1][key] = doc["phases"][0][path][0][key]

            self.expect_reject(edit, needle)

    def test_unknown_enums_rejected(self):
        self.expect_reject(lambda d: d["phases"][0]["tensors"][0].update(kind="embedding"), "张量类型")
        self.expect_reject(lambda d: d["phases"][0]["operators"][0].update(analysis_status="partial"), "分析状态")
        self.expect_reject(lambda d: d["phases"][0]["parameters"][0]["source_refs"][0].update(kind="guess"), "来源类型")

    def test_unresolvable_tensor_references_rejected(self):
        for field in ("inputs", "outputs", "weights"):
            def edit(doc, field=field):
                doc["phases"][0]["operators"][1][field] = ["t_does_not_exist"]

            self.expect_reject(edit, field)

    def test_unresolvable_parameter_refs_rejected(self):
        def op_edit(doc):
            doc["phases"][0]["operators"][1]["parameter_refs"] = ["p_missing"]

        self.expect_reject(op_edit, "算子 op_router.parameter_refs")

        def tensor_edit(doc):
            doc["phases"][0]["tensors"][0]["parameter_refs"] = ["p_missing_too"]

        self.expect_reject(tensor_edit, "张量 t_hidden.parameter_refs")

    def test_phase_scoped_refs_do_not_leak_across_phases(self):
        # p_topk 只存在于 prefill；decode 的算子引用它必须被拒
        def edit(doc):
            doc["phases"][1]["operators"][0]["parameter_refs"] = ["p_topk"]

        self.expect_reject(edit, "p_topk")

    def test_shape_and_type_validation(self):
        def float_axis(doc):
            doc["phases"][0]["tensors"][0]["shape"] = ["B", 1.5, "H"]

        self.expect_reject(float_axis, "shape 轴")

        def not_list(doc):
            doc["phases"][0]["tensors"][0]["shape"] = "B,T,H"

        self.expect_reject(not_list, "shape")

        def empty_symbol(doc):
            doc["phases"][0]["tensors"][0]["shape"] = ["B", "  ", "H"]

        self.expect_reject(empty_symbol, "不能为空")

        def bad_layer(doc):
            doc["scope"]["layer_index"] = "1"

        self.expect_reject(bad_layer, "layer_index")

        def negative_layer(doc):
            doc["scope"]["layer_index"] = -1

        self.expect_reject(negative_layer, "layer_index")

    def test_fixed_enums_of_contract(self):
        self.expect_reject(lambda d: d["model"].update(implementation="gguf"), "implementation")
        self.expect_reject(lambda d: d["scope"].update(analysis_path="graph"), "analysis_path")

    def test_non_json_values_and_control_chars_rejected(self):
        def non_json(doc):
            doc["phases"][0]["parameters"][0]["value"] = {1, 2}

        self.expect_reject(non_json, "JSON")

        def control_char(doc):
            doc["phases"][0]["operators"][0]["note"] = "内部\x00字符"

        self.expect_reject(control_char, "非法控制字符")


class TestEmptyButLegalDocuments(unittest.TestCase):
    def test_empty_sections_are_rendered_without_inventing_rows(self):
        doc = load_doc()
        phase = doc["phases"][0]
        phase["parameters"], phase["tensors"], phase["operators"] = [], [], []
        phase["cache_notes"], phase["coverage_notes"] = [], []
        payload = export_operator_workbook(doc)
        wb, ws = sheet_of(payload, "Prefill")
        self.assertEqual(len(op_data_rows(ws)), 1)  # 仅一条“未提供”说明行
        self.assertEqual(len(param_data_rows(ws)), 1)
        self.assertIn("本阶段文档未提供", cell(ws, op_data_rows(ws)[0], "marker"))
        self.assertIn("本阶段文档未提供", param_cell(ws, param_data_rows(ws)[0], "id"))
        self.assertEqual(wb.sheetnames, ["Prefill", "Decode"])

    def test_phase_ids_must_be_the_frozen_two(self):
        doc = load_doc()
        doc["phases"][0]["id"] = "prefillish"
        with self.assertRaises(ValueError):
            export_operator_workbook(doc)


class TestDuplicateSymbolBinding(unittest.TestCase):
    """修复01回归：同名 symbol 不得静默选错参数（接口只保证 id 阶段内唯一）。

    均在自有合成文档副本上变异，不读 Codex 证据目录、不依赖 coordination 产物。
    """

    def setUp(self):
        self.doc = load_doc()

    def collide(self, phase_index=0, param_id="p_topk", symbol="H"):
        """把某参数的 symbol 改成与 p_hidden 同名（值不动），复现同名候选。"""
        for param in self.doc["phases"][phase_index]["parameters"]:
            if param["id"] == param_id:
                param["symbol"] = symbol
                return
        raise KeyError(param_id)

    def sheet(self, name):
        payload = export_operator_workbook(self.doc)
        return sheet_of(payload, name)[1]

    def rows_by_op(self, ws):
        return {cell(ws, r, "op_id"): r for r in op_data_rows(ws)}

    def candidate_rows(self, ws, pid):
        return next(r for r in param_data_rows(ws) if param_cell(ws, r, "id") == pid)

    def test_explicit_ref_disambiguates_same_symbol(self):
        """本复现：维度符号 H 有两个同名候选时，取本行显式引用的那个参数。"""
        self.collide()                       # p_topk.symbol -> H（值仍 8），p_hidden 仍 H=4096
        ws = self.sheet("Prefill")
        rows = self.rows_by_op(ws)
        topk_row, hidden_row = self.candidate_rows(ws, "p_topk"), self.candidate_rows(ws, "p_hidden")
        r = rows["op_embed_lookup"]          # 只通过 t_hidden 显式引用 p_hidden
        self.assertIn("p_hidden", cell(ws, r, "param_ids"))
        self.assertNotIn("p_topk", cell(ws, r, "param_ids").split(", "))
        axis = cell(ws, r, "axis_lookup")
        self.assertTrue(axis.startswith("="), axis)
        targets = set(formula_targets(axis))
        self.assertIn(("D", str(hidden_row)), targets)   # H 绑到 p_hidden=4096
        self.assertNotIn(("D", str(topk_row)), targets)  # 不再悄悄绑 p_topk=8
        self.assertNotIn("歧义", axis)                    # 能辨认就不报歧义
        # 同行两列必须一致：维度联查取值参数⊆参数联查取值参数
        lookup = set(formula_targets(cell(ws, r, "param_lookup")))
        self.assertTrue({t for t in targets if t[0] == "D"} <= {t for t in lookup if t[0] == "D"})

    def test_row_really_referencing_two_same_symbol_params_is_not_guessed(self):
        """同一行真实引用两个同名参数：不得选一，必须标明歧义并给出可定位id。"""
        self.collide()
        op = next(o for o in self.doc["phases"][0]["operators"] if o["id"] == "op_embed_lookup")
        op["parameter_refs"] = ["p_topk"]    # 与张量带的 p_hidden 同名不同值
        ws = self.sheet("Prefill")
        rows = self.rows_by_op(ws)
        topk_row, hidden_row = self.candidate_rows(ws, "p_topk"), self.candidate_rows(ws, "p_hidden")
        r = rows["op_embed_lookup"]
        axis = cell(ws, r, "axis_lookup")
        self.assertIn("符号存在歧义", axis)
        self.assertIn(f"p_topk（第{topk_row}行）", axis)
        self.assertIn(f"p_hidden（第{hidden_row}行）", axis)
        targets = set(formula_targets(axis))
        for col in ("B", "D"):                # 歧义符号绝不取其中任何一个的值
            self.assertNotIn((col, str(topk_row)), targets)
            self.assertNotIn((col, str(hidden_row)), targets)
        # 参数联查列：同名的两项附id（$A$），避免两个 H 相互混淆
        lookup = cell(ws, r, "param_lookup")
        self.assertIn(("A", str(topk_row)), formula_targets(lookup))
        self.assertIn(("A", str(hidden_row)), formula_targets(lookup))
        self.assertIn(("D", str(topk_row)), formula_targets(lookup))
        self.assertIn(("D", str(hidden_row)), formula_targets(lookup))

    def test_duplicate_symbol_without_any_explicit_ref_is_ambiguous(self):
        """显式引用分辨不出（零候选被引用）：同样不得猜值，但无冲突的符号照旧绑定。"""
        self.collide()
        for tensor in self.doc["phases"][0]["tensors"]:
            tensor["parameter_refs"] = []
        for op in self.doc["phases"][0]["operators"]:
            op["parameter_refs"] = []
        ws = self.sheet("Prefill")
        rows = self.rows_by_op(ws)
        topk_row = self.candidate_rows(ws, "p_topk")
        hidden_row = self.candidate_rows(ws, "p_hidden")
        batch_row = self.candidate_rows(ws, "p_batch")
        r = rows["op_router"]
        axis = cell(ws, r, "axis_lookup")
        self.assertIn("符号存在歧义", axis)
        self.assertIn("未命中本表参数区同名符号：E", axis)
        targets = set(formula_targets(axis))
        self.assertNotIn(("D", str(topk_row)), targets)   # 不猜 p_topk=8
        self.assertNotIn(("D", str(hidden_row)), targets)  # 也不猜 p_hidden=4096
        self.assertIn(("D", str(batch_row)), targets)      # 唯一符号 B 照旧绑定
        # 引用集全空：参数引用列写无引用且不带跳转链接
        self.assertEqual(cell(ws, r, "param_ids"), "（无参数引用）")
        self.assertEqual(cell(ws, r, "param_lookup"), "（无参数引用）")
        self.assertIsNone(ws.cell(row=r, column=COL["param_ids"]).hyperlink)

    def test_same_symbol_in_two_phases_stays_phase_scoped(self):
        """不同 phase 各自同名是合法的：两表只拿本表参数分辨/报歧义。"""
        self.collide(phase_index=0, param_id="p_topk", symbol="H")
        self.collide(phase_index=1, param_id="p_step", symbol="H")
        payload = export_operator_workbook(self.doc)
        wb = load_workbook(io.BytesIO(payload))
        self.assertEqual(wb.sheetnames, ["Prefill", "Decode"])
        wsf, wsd = wb["Prefill"], wb["Decode"]
        f_axis = cell(wsf, self.rows_by_op(wsf)["op_router"], "axis_lookup")
        d_axis = cell(wsd, self.rows_by_op(wsd)["op_attn_step"], "axis_lookup")
        self.assertIn("符号存在歧义", f_axis)
        self.assertIn("符号存在歧义", d_axis)
        # 候选id 只属于本表：prefill 列 p_topk/p_hidden，decode 列 p_step/p_hidden
        self.assertIn("p_topk", f_axis)
        self.assertNotIn("p_step", f_axis)
        self.assertIn("p_step", d_axis)
        self.assertNotIn("p_topk", d_axis)
        for ws, name in ((wsf, "Prefill"), (wsd, "Decode")):
            prow = param_data_rows(ws)
            for r in op_data_rows(ws):
                for _c, row_ref in formula_targets(cell(ws, r, "axis_lookup")):
                    self.assertIn(int(row_ref), prow, (name, r))


def _phase(name: str) -> dict:
    """从合成文档按 sheet 名取对应阶段（测试侧独立解析，不复用被测实现的内部结构）。"""
    key = {"Prefill": "prefill", "Decode": "decode"}[name] if name in ("Prefill", "Decode") else name
    for phase in load_doc()["phases"]:
        if phase["id"] == key:
            return phase
    raise KeyError(name)


if __name__ == "__main__":
    unittest.main(verbosity=2)
