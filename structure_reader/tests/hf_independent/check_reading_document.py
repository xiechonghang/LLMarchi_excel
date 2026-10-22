"""Independent checks that consume an actual hf_decoder_reading.v1 document.

Unlike ``test_hf_*.py`` (source expectations), this module inspects a produced
reading document / HTML / xlsx as the object under test. The expected semantics
are those independently derived from the pinned HF sources (see
``INDEPENDENT_CHECKLIST.md``); the document is never used to derive the
expectation.

Public API:
    check_reading_document(document) -> list[Issue]
    check_html(html_text, document) -> list[Issue]
    check_workbook(path, document) -> list[Issue]
    summarize(issues) -> dict

Issue severities:
    error    — a definite structural/semantic inconsistency with the HF source
    warning  — a suspicious connection/coverage gap that needs review
    manual   — cannot be auto-decided here; requires human role mapping / review

No torch / transformers import; no model execution; no reads of coordination
state or generated figures other than the explicit object under test.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, asdict

SCHEMA = "hf_decoder_reading.v1"
PHASE_ORDER = ("prefill", "decode")
KINDS = {"activation", "weight", "cache", "mask", "index", "boundary"}
STATUS = {"analyzed", "internal_unanalyzed"}
EXPECTED_LAYER_INDEX = 4
EXPECTED_ATTENTION_TYPE = "compressed_sparse_attention"
EXPECTED_MOE_TYPE = "moe"
EXPECTED_ANALYSIS_PATH = "eager"


@dataclass
class Issue:
    severity: str
    code: str
    message: str
    phase: str | None = None
    subject: str | None = None

    def as_dict(self) -> dict:
        return asdict(self)


def _err(code, message, phase=None, subject=None):
    return Issue("error", code, message, phase, subject)


def _warn(code, message, phase=None, subject=None):
    return Issue("warning", code, message, phase, subject)


def _manual(code, message, phase=None, subject=None):
    return Issue("manual", code, message, phase, subject)


def _ops_with(phase, token):
    return [o for o in phase["operators"] if token in (o.get("op") or "")]


def _phase_inputs(phase):
    """External layer inputs are the inputs of the first operator in reading order."""
    ops = phase.get("operators") or []
    return set(ops[0]["inputs"]) if ops else set()


def _phase_outputs(phase):
    ops = phase.get("operators") or []
    return set(ops[-1]["outputs"]) if ops else set()


# ---------------------------------------------------------------------------
# Generic integrity checks (product-name independent)
# ---------------------------------------------------------------------------

def check_document_header(document):
    issues = []
    if document.get("schema") != SCHEMA:
        issues.append(_err("schema", f"schema must be {SCHEMA!r}, got {document.get('schema')!r}"))
    model = document.get("model") or {}
    for field in ("model_key", "display_name", "implementation", "source_revision", "config_path"):
        if not model.get(field):
            issues.append(_err("model_field", f"model.{field} missing/empty"))
    scope = document.get("scope") or {}
    if scope.get("layer_index") != EXPECTED_LAYER_INDEX:
        issues.append(_err("scope_layer", f"scope.layer_index must be {EXPECTED_LAYER_INDEX}, got {scope.get('layer_index')!r}"))
    if scope.get("attention_type") != EXPECTED_ATTENTION_TYPE:
        issues.append(_err("scope_attn", f"scope.attention_type must be {EXPECTED_ATTENTION_TYPE!r}, got {scope.get('attention_type')!r}"))
    if scope.get("moe_type") != EXPECTED_MOE_TYPE:
        issues.append(_err("scope_moe", f"scope.moe_type must be {EXPECTED_MOE_TYPE!r}, got {scope.get('moe_type')!r}"))
    if scope.get("analysis_path") != EXPECTED_ANALYSIS_PATH:
        issues.append(_err("scope_path", f"scope.analysis_path must be {EXPECTED_ANALYSIS_PATH!r}, got {scope.get('analysis_path')!r}"))
    phases = document.get("phases") or []
    ids = [p.get("id") for p in phases]
    if ids != list(PHASE_ORDER):
        issues.append(_err("phase_order", f"phases must be exactly {list(PHASE_ORDER)}, got {ids!r}"))
    return issues


def check_ids(phase):
    issues = []
    pid = phase.get("id")
    for key, label in (("parameters", "parameter"), ("tensors", "tensor"), ("operators", "operator")):
        seen = set()
        for item in phase.get(key, []) or []:
            i = item.get("id")
            if i in seen:
                issues.append(_err("duplicate_id", f"duplicate {label} id {i!r}", pid, i))
            seen.add(i)
    return issues


def check_references(phase):
    issues = []
    pid = phase.get("id")
    params = {p["id"] for p in phase.get("parameters", []) or []}
    tensors = {t["id"] for t in phase.get("tensors", []) or []}
    for t in phase.get("tensors", []) or []:
        for pref in t.get("parameter_refs", []) or []:
            if pref not in params:
                issues.append(_err("dangling_param_ref", f"tensor {t['id']!r} parameter_ref {pref!r} not in phase.parameters", pid, t["id"]))
        if t.get("kind") not in KINDS:
            issues.append(_err("bad_kind", f"tensor {t['id']!r} kind {t.get('kind')!r} not in {sorted(KINDS)}", pid, t["id"]))
        if not isinstance(t.get("shape"), list):
            issues.append(_err("bad_shape", f"tensor {t['id']!r} shape must be a list", pid, t["id"]))
    for o in phase.get("operators", []) or []:
        for key in ("inputs", "outputs", "weights"):
            for tid in o.get(key, []) or []:
                if tid not in tensors:
                    issues.append(_err("dangling_tensor_ref", f"operator {o['id']!r} {key} tensor {tid!r} not in phase.tensors", pid, o["id"]))
        for pref in o.get("parameter_refs", []) or []:
            if pref not in params:
                issues.append(_err("dangling_param_ref", f"operator {o['id']!r} parameter_ref {pref!r} not in phase.parameters", pid, o["id"]))
        if o.get("analysis_status") not in STATUS:
            issues.append(_err("bad_status", f"operator {o['id']!r} analysis_status {o.get('analysis_status')!r} invalid", pid, o["id"]))
    return issues


def _producer_map(phase):
    m = {}
    for idx, o in enumerate(phase.get("operators", []) or []):
        for tid in o.get("outputs", []) or []:
            m.setdefault(tid, []).append(idx)
    return m


def check_single_producer(phase):
    """A tensor may be written in place by several successive ops (later writer consumes
    it), but must not be re-derived by a second op that does not read it first."""
    issues = []
    pid = phase.get("id")
    ops = phase.get("operators", []) or []
    for tid, idxs in _producer_map(phase).items():
        if len(idxs) <= 1:
            continue
        writers = [ops[i] for i in idxs]
        extra = writers[1:]
        all_inplace = all(tid in (w.get("inputs") or []) for w in extra)
        if all_inplace:
            continue  # genuine in-place refinement chain
        names = [w["id"] for w in writers]
        issues.append(_err("duplicate_producer", f"tensor {tid!r} re-derived by {len(idxs)} operators {names} without reading the prior value", pid, tid))
    return issues


def check_producer_presence(phase):
    issues = []
    pid = phase.get("id")
    tmap = {t["id"]: t for t in phase.get("tensors", []) or []}
    produced = set(_producer_map(phase))
    inputs = _phase_inputs(phase)
    consumed = set()
    for o in phase.get("operators", []) or []:
        consumed.update(o.get("inputs", []) or [])
    for tid in sorted(consumed):
        t = tmap.get(tid)
        if t is None:
            continue
        if t.get("kind") in ("boundary", "cache", "weight"):
            continue
        if tid in inputs or tid in produced:
            continue
        issues.append(_err("missing_producer", f"consumed {t.get('kind')} tensor {tid!r} has no producer and is not a declared phase input", pid, tid))
    return issues


def check_topological_order(phase):
    issues = []
    pid = phase.get("id")
    pm = _producer_map(phase)
    tmap = {t["id"]: t for t in phase.get("tensors", []) or []}
    for idx, o in enumerate(phase.get("operators", []) or []):
        outs = set(o.get("outputs", []) or [])
        for tid in o.get("inputs", []) or []:
            if tid in outs:
                continue  # in-place read/write of the same tensor
            t = tmap.get(tid)
            if t is not None and t.get("kind") == "cache":
                continue  # reading prior cache state produced by a later op is fine
            prod = pm.get(tid)
            if prod and min(prod) > idx:
                pidx = min(prod)
                issues.append(_err(
                    "ordering",
                    f"operator {o['id']!r} (idx {idx}) consumes {tid!r} produced later by {phase['operators'][pidx]['id']!r} (idx {pidx})",
                    pid, o["id"],
                ))
    return issues


def check_unused_and_unreferenced(phase):
    issues = []
    pid = phase.get("id")
    tensors = phase.get("tensors", []) or []
    ops = phase.get("operators", []) or []
    referenced = set()
    for o in ops:
        referenced.update(o.get("inputs", []) or [])
        referenced.update(o.get("outputs", []) or [])
        referenced.update(o.get("weights", []) or [])
    consumed = set()
    for o in ops:
        consumed.update(o.get("inputs", []) or [])
    # A decoder layer preserves the hidden-stream shape, so any activation tensor that
    # shares the phase input shape is a passthrough (phase) output, not a dead end.
    input_shapes = {tuple(t["shape"]) for t in tensors if t["id"] in _phase_inputs(phase) and isinstance(t.get("shape"), list)}
    passthrough = {t["id"] for t in tensors if t.get("kind") == "activation" and isinstance(t.get("shape"), list) and tuple(t["shape"]) in input_shapes}
    for t in tensors:
        tid, kind = t["id"], t.get("kind")
        if tid not in referenced:
            if kind == "boundary":
                continue  # boundary may be declared even if this layer does not consume it
            issues.append(_warn("unreferenced_tensor", f"tensor {tid!r} (kind={kind}) is declared but never referenced by any operator", pid, tid))
            continue
        if kind == "activation" and tid not in consumed and tid not in passthrough:
            issues.append(_warn("unused_output", f"activation tensor {tid!r} is produced but never consumed and is not a phase output", pid, tid))
    return issues


# ---------------------------------------------------------------------------
# Semantic checks with independently derived expectations
# ---------------------------------------------------------------------------

def _resolve_attention_ops(phase):
    """Resolve the eager attention softmax / AV / repeat-KV ops by structure, not by note.

    AV is the unique op doing matmul+transpose (no softmax); its two activation inputs
    are produced by the attention softmax and by the KV head-broadcast (expand/repeat).
    Returns (av, softmax, repeat, kv_rep) or (av, None, ...) when unresolvable.
    """
    ops = phase.get("operators", []) or []
    tmap = {t["id"]: t for t in phase.get("tensors", []) or []}
    pm = _producer_map(phase)
    av_ops = [o for o in ops if "matmul" in (o.get("op") or "") and "transpose" in (o.get("op") or "") and "softmax" not in (o.get("op") or "")]
    if len(av_ops) != 1:
        return (None, None, None, None)
    av = av_ops[0]
    softmax = repeat = None
    kv_rep = None
    for tid in av.get("inputs", []) or []:
        if tmap.get(tid, {}).get("kind") == "weight":
            continue
        for pi in pm.get(tid, []):
            pop = ops[pi]
            if "softmax" in (pop.get("op") or ""):
                softmax = pop
            if "expand" in (pop.get("op") or "") or "repeat" in (pop.get("id") or "") or "broadcast" in (pop.get("op") or ""):
                repeat = pop
                kv_rep = tid
    return (av, softmax, repeat, kv_rep)


def check_q_residual_norm_weight(phase):
    """The tensor feeding the Q up-projection must be produced by a *weighted* RMSNorm
    (source: q_residual = q_a_norm(q_a_proj(hidden)); q_a_norm is DeepseekV4RMSNorm)."""
    issues = []
    pid = phase.get("id")
    tmap = {t["id"]: t for t in phase.get("tensors", []) or []}
    pm = _producer_map(phase)
    ops = phase.get("operators", []) or []
    q_up = [o for o in ops if "linear" in (o.get("op") or "") and "transpose" in (o.get("op") or "") and "matmul" not in (o.get("op") or "")]
    if len(q_up) != 1:
        issues.append(_manual("q_residual_norm_weight", f"could not uniquely resolve the Q up-projection op (candidates={[o['id'] for o in q_up]})", pid))
        return issues
    q_up_op = q_up[0]
    if len(q_up_op.get("inputs") or []) != 1:
        issues.append(_manual("q_residual_norm_weight", f"Q up-projection {q_up_op['id']!r} has {len(q_up_op.get('inputs') or [])} inputs", pid, q_up_op["id"]))
        return issues
    q_in = q_up_op["inputs"][0]
    producers = [ops[i] for i in pm.get(q_in, [])]
    weighted_norm = None
    for p in producers:
        if "rmsnorm" not in (p.get("op") or "").lower():
            continue
        has_weight = bool(p.get("weights")) or any(tmap.get(i, {}).get("kind") == "weight" for i in (p.get("inputs") or []))
        weighted_norm = has_weight
    if weighted_norm is None:
        issues.append(_err("q_residual_norm_weight", f"Q up-projection consumes {q_in!r} whose producer is not an RMSNorm op (source: q_residual is RMS-normalized)", pid, q_up_op["id"]))
    elif not weighted_norm:
        issues.append(_err("q_residual_norm_weight", f"Q residual RMSNorm (producer of {q_in!r}) carries no learned weight (source: DeepseekV4RMSNorm has a weight)", pid, q_up_op["id"]))
    return issues


def check_attention_prob_axis(phase):
    """The eager attention probability key axis must equal the KV sequence length read by AV.

    Source: probs over concatenated KV of length L_kv (incl. compressed entries); the
    declared softmax output axis must match the repeated-KV sequence axis.
    """
    issues = []
    pid = phase.get("id")
    tmap = {t["id"]: t for t in phase.get("tensors", []) or []}
    av, softmax, repeat, kv_rep = _resolve_attention_ops(phase)
    if av is None or softmax is None or kv_rep is None:
        issues.append(_manual("attn_prob_axis", "could not resolve the attention softmax/AV/repeat-KV chain", pid))
        return issues
    probs = softmax["outputs"][0]
    pshape, kshape = tmap[probs]["shape"], tmap[kv_rep]["shape"]
    if not pshape or not kshape:
        return issues
    key_axis = pshape[-1]
    kv_seq_axis = kshape[-2] if len(kshape) >= 3 else kshape[-1]
    if key_axis != kv_seq_axis:
        issues.append(_err(
            "attn_prob_axis",
            f"softmax output {probs!r} key axis {key_axis!r} != repeated-KV seq axis {kv_seq_axis!r} (source: probabilities span the full KV length L_kv+compressed)",
            pid, probs,
        ))
    return issues


def check_av_inputs(phase):
    """AV must consume exactly the softmax probabilities and the repeated KV, not q/qk logits."""
    issues = []
    pid = phase.get("id")
    tmap = {t["id"]: t for t in phase.get("tensors", []) or []}
    av, softmax, repeat, kv_rep = _resolve_attention_ops(phase)
    if av is None or softmax is None:
        issues.append(_manual("av_inputs", "could not resolve the attention softmax/AV chain", pid))
        return issues
    probs = softmax["outputs"][0]
    act_inputs = [i for i in (av.get("inputs") or []) if tmap.get(i, {}).get("kind") != "weight"]
    expected = {probs, kv_rep}
    if kv_rep is None or set(act_inputs) != expected:
        issues.append(_err("av_inputs", f"AV op {av['id']!r} activation inputs {act_inputs} != expected softmax={probs!r} + repeated KV={kv_rep!r}", pid, av["id"]))
    return issues


def check_mhc_transpose_chain(phase):
    """Each single-input transpose op must keep rank with the last two axes swapped,
    and its result must be consumed by the residual mix (comb.T @ residual)."""
    issues = []
    pid = phase.get("id")
    tmap = {t["id"]: t for t in phase.get("tensors", []) or []}
    consumed = set()
    for o in phase.get("operators", []) or []:
        consumed.update(o.get("inputs", []) or [])
    found = 0
    for o in phase.get("operators", []) or []:
        op = o.get("op") or ""
        if "transpose" not in op:
            continue
        # only a *pure* transpose counts; linear+transpose / matmul+transpose are other ops
        if any(tok in op for tok in ("linear", "matmul", "bmm", "reshape", "view")):
            continue
        if len(o.get("inputs", []) or []) != 1 or len(o.get("outputs", []) or []) != 1:
            continue
        found += 1
        src, dst = o["inputs"][0], o["outputs"][0]
        sshape, dshape = tmap.get(src, {}).get("shape") or [], tmap.get(dst, {}).get("shape") or []
        expected = list(sshape)
        if len(expected) >= 2:
            expected[-2], expected[-1] = expected[-1], expected[-2]
        if list(dshape) != expected:
            issues.append(_err("transpose_rank", f"transpose op {o['id']!r}: {src}{sshape} -> {dst}{dshape}, expected {expected} (transpose swaps the last two axes)", pid, o["id"]))
        if dst not in consumed:
            issues.append(_err("transpose_isolated", f"transpose op {o['id']!r} output {dst!r} is never consumed; the residual mix must use comb.T (source: torch.matmul(comb.transpose(-1,-2), hidden))", pid, o["id"]))
    if found == 0:
        issues.append(_manual("transpose_chain", "no single-input transpose op found for the mHC comb.T chain", pid))
    return issues


def check_indexer_query_norm(phase):
    """Indexer query projection must consume the normalized Q residual (q_a_norm output),
    not the raw low-rank Q (source: q_residual = q_a_norm(q_a_proj(...)))."""
    issues = []
    pid = phase.get("id")
    tmap = {t["id"]: t for t in phase.get("tensors", []) or []}
    pm = _producer_map(phase)
    ops = phase.get("operators", []) or []
    # indexer query projection: linear+rope op with a weight whose first axis mentions index heads
    cand = []
    for o in ops:
        op = o.get("op") or ""
        if "linear" in op and "rope" in op and "transpose" not in op:
            wsh = [tmap.get(w, {}).get("shape") or [] for w in (o.get("weights") or [])]
            cand.append((o, wsh))
    if len(cand) != 1:
        issues.append(_manual("indexer_query_norm", f"could not uniquely resolve the indexer query projection op (candidates={[c[0]['id'] for c in cand]})", pid))
        return issues
    o, _ = cand[0]
    ins = o.get("inputs") or []
    if len(ins) != 1:
        issues.append(_manual("indexer_query_norm", f"indexer query op {o['id']!r} has {len(ins)} inputs; cannot verify normalized-Q identity", pid, o["id"]))
        return issues
    q_in = ins[0]
    producers = pm.get(q_in) or []
    producer_ops = [ops[i] for i in producers]
    is_weighted_norm = any("rmsnorm" in (p.get("op") or "").lower() and "unweight" not in (p.get("op") or "").lower() for p in producer_ops)
    if not is_weighted_norm:
        issues.append(_err(
            "indexer_query_norm",
            f"indexer query op {o['id']!r} consumes {q_in!r}, not the output of a weighted RMSNorm on the Q residual (source: indexer consumes q_residual = q_a_norm(q_a_proj(hidden)))",
            pid, o["id"],
        ))
    return issues


def check_expert_input_linkage(phase):
    """The routed expert compute op must consume an external routed-token activation,
    not only tensors it produces itself (source: F.linear(hidden_states[token_idx], gate_up))."""
    issues = []
    pid = phase.get("id")
    tmap = {t["id"]: t for t in phase.get("tensors", []) or []}
    ops = phase.get("operators", []) or []
    expert_ops = [o for o in ops if any(len(tmap.get(w, {}).get("shape") or []) == 3 for w in (o.get("weights") or []))]
    if not expert_ops:
        issues.append(_manual("expert_input_linkage", "could not locate a routed expert op with 3-D expert weights", pid))
        return issues
    for o in expert_ops:
        outs = set(o.get("outputs") or [])
        act_inputs = [i for i in (o.get("inputs") or []) if tmap.get(i, {}).get("kind") != "weight"]
        external = [i for i in act_inputs if i not in outs]
        if not external:
            issues.append(_err(
                "expert_input_linkage",
                f"expert op {o['id']!r} consumes no routed-token activation distinct from its own outputs {sorted(outs)} (source: expert compute reads the dispatched hidden states)",
                pid, o["id"],
            ))
    return issues


def check_phase_notes(phase):
    issues = []
    pid = phase.get("id")
    for key in ("cache_notes", "coverage_notes"):
        if not isinstance(phase.get(key), list) or not phase.get(key):
            issues.append(_err("missing_notes", f"phase {pid!r} must carry a non-empty {key} list", pid))
    return issues


def check_reading_document(document):
    issues = list(check_document_header(document))
    for phase in document.get("phases", []) or []:
        issues += check_ids(phase)
        issues += check_references(phase)
        issues += check_single_producer(phase)
        issues += check_producer_presence(phase)
        issues += check_topological_order(phase)
        issues += check_unused_and_unreferenced(phase)
        issues += check_q_residual_norm_weight(phase)
        issues += check_attention_prob_axis(phase)
        issues += check_av_inputs(phase)
        issues += check_mhc_transpose_chain(phase)
        issues += check_indexer_query_norm(phase)
        issues += check_expert_input_linkage(phase)
        issues += check_phase_notes(phase)
    return issues


# ---------------------------------------------------------------------------
# HTML / xlsx checks
# ---------------------------------------------------------------------------

def check_html(html_text, document):
    issues = []
    ids = set(re.findall(r'\bid="([^"]+)"', html_text))
    hrefs = re.findall(r'href="#([^"]+)"', html_text)
    missing = sorted(set(hrefs) - ids)
    if missing:
        issues.append(_err("html_dangling_href", f"{len(hrefs)} in-page links reference {len(missing)} missing targets, e.g. {missing[:8]}"))
    if 'id="phase-prefill"' in html_text and 'id="phase-decode"' in html_text:
        if html_text.index('id="phase-prefill"') > html_text.index('id="phase-decode"'):
            issues.append(_err("html_phase_order", "decode phase appears before prefill phase in the document"))
    else:
        issues.append(_err("html_phase_sections", 'missing one of id="phase-prefill" / id="phase-decode"'))
    if len(re.findall(r'<section[^>]*class="[^"]*\bphase\b', html_text)) not in (0, 2):
        issues.append(_warn("html_phase_count", "expected exactly two phase sections"))
    has_blackbox = any(
        o.get("analysis_status") == "internal_unanalyzed"
        for p in (document.get("phases") or [])
        for o in (p.get("operators") or [])
    )
    if has_blackbox and "内部未分析" not in html_text:
        issues.append(_err("html_blackbox_marker", "document contains internal_unanalyzed ops but HTML lacks the 内部未分析 marker"))
    external = re.findall(r'(?:src|href)="(https?://[^"]+)"', html_text)
    if external:
        issues.append(_warn("html_external_ref", f"{len(external)} external http(s) resource reference(s) found; offline delivery expected: {external[:5]}"))
    return issues


def _find_header(ws, token):
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and token in v:
                return r, c
    return None, None


def _header_cols(ws, header_row):
    return {ws.cell(header_row, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(header_row, c).value}


def check_workbook(path, document):
    """Read-only xlsx check: two sheets, same-source operator/param identity, resolvable
    in-sheet parameter references, and shape text matching the JSON tensors."""
    issues = []
    try:
        import openpyxl
    except Exception as exc:  # pragma: no cover - environment dependent
        return [_manual("xlsx_openpyxl", f"openpyxl unavailable: {exc}")]
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        return [_err("xlsx_unreadable", f"cannot read workbook: {exc}")]
    if wb.sheetnames != ["Prefill", "Decode"]:
        issues.append(_err("xlsx_sheets", f"sheets must be exactly ['Prefill','Decode'], got {wb.sheetnames}"))
    doc_phases = {p["id"]: p for p in document.get("phases", []) or []}
    for pid, sheet_name in (("prefill", "Prefill"), ("decode", "Decode")):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        doc = doc_phases.get(pid, {})
        pr, pc = _find_header(ws, "参数id")
        orow, oidc = _find_header(ws, "算子id")
        if pr is None or orow is None:
            issues.append(_err("xlsx_regions", f"{sheet_name}: missing 参数id or 算子id header", pid))
            continue
        pcols = _header_cols(ws, pr)
        ocols = _header_cols(ws, orow)
        # collect param ids
        param_ids = []
        r = pr + 1
        while r <= ws.max_row:
            v = ws.cell(r, pc).value
            if v is None or str(v).strip() == "":
                break
            param_ids.append(str(v).strip())
            r += 1
        doc_params = [p["id"] for p in doc.get("parameters", []) or []]
        if set(param_ids) != set(doc_params):
            issues.append(_err("xlsx_param_ids", f"{sheet_name}: param id set differs from JSON (sheet {len(param_ids)} vs doc {len(doc_params)})", pid))
        # collect operators
        def col_of(token):
            for k, c in ocols.items():
                if isinstance(k, str) and token in k:
                    return c
            return None
        cin, cout, cw, cref, cname = col_of("输入"), col_of("输出"), col_of("权重"), col_of("参数引用"), col_of("名称")
        tmap = {t["id"]: t for t in doc.get("tensors", []) or []}
        doc_ops = {o["id"]: o for o in doc.get("operators", []) or []}
        sheet_ops = {}
        r = orow + 1
        while r <= ws.max_row:
            v = ws.cell(r, oidc).value
            if v is None or str(v).strip() == "":
                break
            sheet_ops[str(v).strip()] = r
            r += 1
        if set(sheet_ops) != set(doc_ops):
            issues.append(_err("xlsx_op_ids", f"{sheet_name}: operator id set differs from JSON (sheet {len(sheet_ops)} vs doc {len(doc_ops)})", pid))
        param_set = set(param_ids)
        shape_mismatch = 0
        param_link_bad = 0
        name_mismatch = 0
        for oid, rr in sheet_ops.items():
            if oid not in doc_ops:
                continue
            o = doc_ops[oid]
            if cname and ws.cell(rr, cname).value is not None:
                if str(ws.cell(rr, cname).value).strip() != str(o.get("name", "")).strip():
                    name_mismatch += 1
            txt = ref_text = None
            if cref:
                ref_text = ws.cell(rr, cref).value
            if ref_text:
                for tok in re.split(r"[,;/、\s]+", str(ref_text).strip()):
                    if tok and tok not in param_set:
                        param_link_bad += 1
            for ids, ci in ((o.get("inputs"), cin), (o.get("outputs"), cout), (o.get("weights"), cw)):
                if ci is None:
                    continue
                cell = ws.cell(rr, ci).value or ""
                for tid in ids or []:
                    sh = tmap.get(tid, {}).get("shape")
                    if sh is None:
                        continue
                    pat = "[" + ", ".join(sh) + "]"
                    if pat not in str(cell):
                        shape_mismatch += 1
        if shape_mismatch:
            issues.append(_err("xlsx_shape_mismatch", f"{sheet_name}: {shape_mismatch} tensor shape(s) from JSON not found in the operator rows", pid))
        if param_link_bad:
            issues.append(_err("xlsx_param_link", f"{sheet_name}: {param_link_bad} parameter reference token(s) do not resolve to a parameter id in this sheet", pid))
        if name_mismatch:
            issues.append(_err("xlsx_op_name", f"{sheet_name}: {name_mismatch} operator name(s) differ from the JSON document", pid))
    return issues


def summarize(issues):
    out = {"total": len(issues), "error": 0, "warning": 0, "manual": 0, "by_code": {}}
    for i in issues:
        out[i.severity] = out.get(i.severity, 0) + 1
        out["by_code"][i.code] = out["by_code"].get(i.code, 0) + 1
    return out


__all__ = [
    "Issue",
    "check_reading_document",
    "check_html",
    "check_workbook",
    "summarize",
    "SCHEMA",
    "PHASE_ORDER",
]
