"""Self-contained tests for the round01 gate checks.

Builds a synthetic layer4 reading document (clearly synthetic, not the real
candidate) that must satisfy every independent expectation, then mutates exactly
one property per negative test. Never reads outputs / coordination / the real
candidate. The workspace-root argument is only used to locate the read-only HF
inputs for identity checks.
"""

import copy
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

import check_layer4 as cl  # noqa: E402
import check_html as ch  # noqa: E402
import check_xlsx as cx  # noqa: E402

WORKSPACE_ROOT = HERE.parents[3]  # structure_reader/tests/hf_independent/gate_round01 -> repo root


def _t(tid, kind, shape, refs=()):
    return {"id": tid, "kind": kind, "name": tid, "shape": list(shape), "parameter_refs": list(refs), "note": ""}


def _o(oid, op, module, inputs, outputs, weights=(), params=()):
    return {
        "id": oid, "module_path": module, "name": oid, "op": op,
        "inputs": list(inputs), "outputs": list(outputs), "weights": list(weights),
        "parameter_refs": list(params), "source_refs": [], "analysis_status": "analyzed", "note": "",
    }


def _phase(pid):
    ops = []
    W = lambda tid: tid  # noqa: E731
    # ---- mHC sites -------------------------------------------------------
    for site, stream_in in (("attn", "streams_in"), ("ffn", "attn_out_streams")):
        m = f"DeepseekV4DecoderLayer.{site}_hc"
        ops += [
            _o(f"{site}_hc_flatten", "view+float+rmsnorm", m, [stream_in], [f"{site}_mhc_normed"]),
            _o(f"{site}_hc_project", "linear", m, [f"{site}_mhc_normed"], [f"{site}_mhc_weights"], weights=[f"{site}_hc_fn"]),
            _o(f"{site}_hc_split", "split+view", m, [f"{site}_mhc_weights"], [f"{site}_mhc_pre_w", f"{site}_mhc_post_w", f"{site}_mhc_comb_w"]),
            _o(f"{site}_hc_pre", "affine+sigmoid+add", m, [f"{site}_mhc_pre_w", f"{site}_hc_base", f"{site}_hc_scale"], [f"{site}_mhc_pre"]),
            _o(f"{site}_hc_post", "affine+sigmoid+mul", m, [f"{site}_mhc_post_w", f"{site}_hc_base", f"{site}_hc_scale"], [f"{site}_mhc_post"]),
            _o(f"{site}_hc_comb", "affine+softmax+sinkhorn", m, [f"{site}_mhc_comb_w", f"{site}_hc_base", f"{site}_hc_scale"], [f"{site}_mhc_comb"]),
            _o(f"{site}_hc_collapse", "broadcast-mul+sum", m, [f"{site}_mhc_pre", stream_in], [f"{site}_mhc_collapsed"]),
        ]
    # ---- attention / q / kv ---------------------------------------------
    ops += [
        _o("q_a_proj", "linear", "DeepseekV4Attention", ["attn_mhc_collapsed"], ["q_a"], weights=["q_weight"]),
        _o("q_a_norm", "rmsnorm", "DeepseekV4Attention.q_a_norm", ["q_a", "q_a_norm_weight"], ["q_a_normed"]),
        _o("q_b_proj", "linear+view+transpose", "DeepseekV4Attention", ["q_a_normed"], ["q_head"], weights=["q_b_weight"]),
        _o("q_b_norm", "rmsnorm-unweighted", "DeepseekV4Attention.q_b_norm", ["q_head"], ["q_normed"]),
        _o("rope_compress", "outer+cos/sin", "DeepseekV4RotaryEmbedding", ["position_ids"], ["cos_q", "sin_q"]),
        _o("q_rope_apply", "rope-interleaved", "apply_rotary_pos_emb", ["q_normed", "cos_q", "sin_q"], ["q_rope"]),
        _o("kv_proj", "linear+rmsnorm", "DeepseekV4Attention", ["attn_mhc_collapsed"], ["kv_proj_out"], weights=["kv_weight", "kv_norm_weight"]),
        _o("kv_head_view", "view+transpose", "DeepseekV4Attention", ["kv_proj_out"], ["kv_head"]),
        _o("kv_rope_apply", "rope-interleaved", "apply_rotary_pos_emb", ["kv_head", "cos_q", "sin_q"], ["kv_rope"]),
        _o("cache_update", "cat+slice", "DeepseekV4CSACache.update", ["kv_rope", "cache_keys"], ["kv_full", "cache_keys_post"]),
    ]
    # ---- CSA compressor --------------------------------------------------
    ops += [
        _o("c_kv_gate_proj", "linear×2", "DeepseekV4CSACompressor", ["attn_mhc_collapsed"], ["c_kv", "c_gate"], weights=["c_kv_weight", "c_gate_weight"]),
        _o("c_store_window", "cat+slice", "DeepseekV4CSACache.store_compression_weights", ["c_kv", "c_gate", "cache_buffer_c_kv", "cache_buffer_c_gate", "cache_entry_count_c"], ["c_chunk_kv", "c_chunk_gate", "cache_buffer_c_kv_post", "cache_buffer_c_gate_post"]),
        _o("c_window_view", "view+assign", "DeepseekV4CSACompressor", ["c_chunk_kv", "c_chunk_gate", "cache_overlap_c_kv", "cache_overlap_c_gate"], ["c_windows_kv", "c_windows_gate", "cache_overlap_c_kv_post", "cache_overlap_c_gate_post"], weights=["c_position_bias"]),
        _o("c_pool", "softmax+mul+sum+rmsnorm", "DeepseekV4CSACompressor", ["c_windows_kv", "c_windows_gate", "c_kv_norm_weight"], ["c_compressed"]),
        _o("c_rope_gen", "arange+outer+cos/sin", "DeepseekV4CSACompressor.rotary_emb", ["cache_entry_count_c"], ["c_cos", "c_sin"]),
        _o("c_rope", "rope-interleaved", "apply_rotary_pos_emb（CSA 压缩条目）", ["c_compressed", "c_cos", "c_sin"], ["c_compressed_rope"]),
        _o("c_update_states", "cat+add", "DeepseekV4CSACache.update_compressor_states", ["c_compressed_rope", "cache_compressed_c", "cache_entry_count_c"], ["cache_compressed_c_post", "cache_entry_count_c_post"]),
    ]
    # ---- Indexer ---------------------------------------------------------
    ops += [
        _o("i_kv_gate_proj", "linear×2", "DeepseekV4Indexer", ["attn_mhc_collapsed"], ["i_kv", "i_gate"], weights=["i_kv_weight", "i_gate_weight"]),
        _o("i_store_window", "cat+slice", "DeepseekV4CSACache.store_compression_weights", ["i_kv", "i_gate", "cache_buffer_i_kv", "cache_buffer_i_gate", "cache_entry_count_i"], ["i_chunk_kv", "i_chunk_gate", "cache_buffer_i_kv_post", "cache_buffer_i_gate_post"]),
        _o("i_window_view", "view+assign", "DeepseekV4Indexer", ["i_chunk_kv", "i_chunk_gate", "cache_overlap_i_kv", "cache_overlap_i_gate"], ["i_windows_kv", "i_windows_gate", "cache_overlap_i_kv_post", "cache_overlap_i_gate_post"], weights=["i_position_bias"]),
        _o("i_pool_rope", "softmax+mul+sum+rmsnorm+rope", "DeepseekV4Indexer", ["i_windows_kv", "i_windows_gate", "i_kv_norm_weight", "cache_entry_count_i"], ["i_compressed"]),
        _o("i_update_states", "cat+add", "DeepseekV4CSACache.update_compressor_states", ["i_compressed", "cache_compressed_i", "cache_entry_count_i"], ["cache_compressed_i_post", "cache_entry_count_i_post"]),
        _o("i_q_proj", "linear+view+rope", "DeepseekV4Indexer", ["q_a_normed", "cos_q", "sin_q"], ["i_q"], weights=["i_q_weight"]),
        _o("i_scorer", "matmul+relu+mul+sum", "DeepseekV4IndexerScorer", ["i_q", "cache_compressed_i_post", "attn_mhc_collapsed", "i_w_weight"], ["i_scores"]),
        _o("i_topk", "mask+topk+where", "DeepseekV4Indexer", ["i_scores", "position_ids", "cache_compressed_i_post"], ["indexer_topk"]),
        _o("c_block_bias", "scatter", "DeepseekV4CSACompressor", ["cache_compressed_c_post", "indexer_topk"], ["c_block_bias"]),
        _o("kv_concat", "cat", "DeepseekV4Attention", ["kv_full", "cache_compressed_c_post"], ["kv_all"]),
        _o("mask_model_build", "mask-materialize", "masking_utils.create_sliding_window_causal_mask", ["cache_cumlen"], ["mask_model"]),
        _o("mask_extend", "cat", "DeepseekV4Attention", ["mask_model", "c_block_bias"], ["mask_ext"]),
        _o("repeat_kv", "expand+reshape", "eager_attention_forward", ["kv_all"], ["kv_repeated"]),
        _o("qk_scaled", "matmul+mul", "eager_attention_forward", ["q_rope", "kv_repeated"], ["qk_logits"]),
        _o("mask_add", "add", "eager_attention_forward", ["qk_logits", "mask_ext"], ["qk_masked"]),
        _o("sink_softmax", "cat+sub+softmax+slice", "eager_attention_forward", ["qk_masked", "sinks"], ["attn_probs"]),
        _o("av_matmul", "matmul+transpose", "eager_attention_forward", ["attn_probs", "kv_repeated"], ["attn_head_out"]),
        _o("attn_inv_rope", "rope-conjugate", "DeepseekV4Attention", ["attn_head_out", "cos_q", "sin_q"], ["attn_rope_inv"]),
        _o("attn_group_reshape", "reshape", "DeepseekV4Attention", ["attn_rope_inv"], ["attn_grouped"]),
        _o("o_a_proj", "reshape+bmm", "DeepseekV4GroupedLinear", ["attn_grouped"], ["o_a_out"], weights=["o_a_weight"]),
        _o("o_b_proj", "linear+flatten", "DeepseekV4Attention", ["o_a_out"], ["attn_block_output"], weights=["o_b_weight"]),
        _o("attn_residual_combT", "transpose", "DeepseekV4DecoderLayer.forward", ["attn_mhc_comb"], ["attn_combT"]),
        _o("attn_residual_expand", "broadcast-mul", "DeepseekV4DecoderLayer.forward", ["attn_mhc_post", "attn_block_output"], ["attn_expanded"]),
        _o("attn_residual_mix", "matmul+add", "DeepseekV4DecoderLayer.forward", ["attn_combT", "streams_in", "attn_expanded"], ["attn_out_streams"]),
    ]
    # ---- MoE -------------------------------------------------------------
    ops += [
        _o("moe_flatten", "view", "DeepseekV4SparseMoeBlock", ["ffn_mhc_collapsed"], ["ffn_flat"]),
        _o("moe_router", "linear+softplus+sqrt+topk+gather+norm", "DeepseekV4TopKRouter", ["ffn_flat", "router_bias"], ["router_logits", "topk_indices", "topk_weights"], weights=["router_weight"], params=["K", "Scl", "E"]),
        _o("moe_dispatch", "one_hot+nonzero+where", "DeepseekV4Experts", ["topk_indices", "ffn_flat"], ["expert_mask", "hit_experts"]),
        _o("moe_select_tokens", "where+index", "DeepseekV4Experts（专家模板）", ["ffn_flat", "expert_mask", "hit_experts"], ["expert_tokens", "expert_token_idx", "expert_topk_pos"]),
        _o("moe_expert_slice", "index", "DeepseekV4Experts（专家模板）", ["expert_gate_up", "expert_down", "hit_experts"], ["expert_gate_up_e", "expert_down_e"]),
        _o("moe_expert_gu", "linear", "DeepseekV4Experts（专家模板）", ["expert_tokens"], ["expert_gu_out"], weights=["expert_gate_up_e"]),
        _o("moe_expert_clamp", "chunk+clamp", "DeepseekV4Experts（专家模板）", ["expert_gu_out"], ["expert_gate_c", "expert_up_c"]),
        _o("moe_expert_act", "silu+mul", "DeepseekV4Experts（专家模板）", ["expert_gate_c", "expert_up_c"], ["expert_gated"]),
        _o("moe_expert_down", "linear", "DeepseekV4Experts（专家模板）", ["expert_gated"], ["expert_down_raw"], weights=["expert_down_e"]),
        _o("moe_expert_weight", "gather+mul", "DeepseekV4Experts（专家模板）", ["expert_down_raw", "topk_weights", "expert_token_idx", "expert_topk_pos"], ["expert_down_out"]),
        _o("moe_combine", "index_add", "DeepseekV4Experts", ["expert_down_out", "expert_token_idx"], ["routed_out"]),
        _o("moe_routed_view", "view", "DeepseekV4SparseMoeBlock", ["routed_out"], ["routed_view"]),
        _o("shared_gate_up", "linear+clamp×2", "DeepseekV4MLP", ["ffn_mhc_collapsed"], ["shared_gate", "shared_up"], weights=["shared_gate_weight", "shared_up_weight"]),
        _o("shared_act_down", "silu+mul+linear", "DeepseekV4MLP", ["shared_gate", "shared_up"], ["shared_act", "shared_out"], weights=["shared_down_weight"]),
        _o("moe_merge", "add", "DeepseekV4SparseMoeBlock", ["routed_view", "shared_out"], ["mlp_output"]),
        _o("ffn_residual_combT", "transpose", "DeepseekV4DecoderLayer.forward", ["ffn_mhc_comb"], ["ffn_combT"]),
        _o("ffn_residual_expand", "broadcast-mul", "DeepseekV4DecoderLayer.forward", ["ffn_mhc_post", "mlp_output"], ["ffn_expanded"]),
        _o("ffn_residual_mix", "matmul+add", "DeepseekV4DecoderLayer.forward", ["ffn_combT", "attn_out_streams", "ffn_expanded"], ["streams_out"]),
    ]
    shapes = {
        "streams_in": ["B", "T", "Nhc", "H"], "attn_out_streams": ["B", "T", "Nhc", "H"], "streams_out": ["B", "T", "Nhc", "H"],
        "attn_mhc_normed": ["B", "T", "Nhc×H"], "ffn_mhc_normed": ["B", "T", "Nhc×H"],
        "attn_mhc_weights": ["B", "T", "(2+Nhc)×Nhc"], "ffn_mhc_weights": ["B", "T", "(2+Nhc)×Nhc"],
        "attn_mhc_pre_w": ["B", "T", "Nhc"], "attn_mhc_post_w": ["B", "T", "Nhc"], "attn_mhc_comb_w": ["B", "T", "Nhc", "Nhc"],
        "ffn_mhc_pre_w": ["B", "T", "Nhc"], "ffn_mhc_post_w": ["B", "T", "Nhc"], "ffn_mhc_comb_w": ["B", "T", "Nhc", "Nhc"],
        "attn_mhc_pre": ["B", "T", "Nhc"], "attn_mhc_post": ["B", "T", "Nhc"], "attn_mhc_comb": ["B", "T", "Nhc", "Nhc"],
        "ffn_mhc_pre": ["B", "T", "Nhc"], "ffn_mhc_post": ["B", "T", "Nhc"], "ffn_mhc_comb": ["B", "T", "Nhc", "Nhc"],
        "attn_mhc_collapsed": ["B", "T", "H"], "ffn_mhc_collapsed": ["B", "T", "H"],
        "q_a": ["B", "T", "Rq"], "q_a_normed": ["B", "T", "Rq"], "q_head": ["B", "Nh", "T", "Dh"], "q_normed": ["B", "Nh", "T", "Dh"], "q_rope": ["B", "Nh", "T", "Dh"],
        "cos_q": ["T", "Rk÷2"], "sin_q": ["T", "Rk÷2"],
        "kv_proj_out": ["B", "T", "Dh"], "kv_head": ["B", "1", "T", "Dh"], "kv_rope": ["B", "1", "T", "Dh"],
        "kv_full": ["B", "1", "L_kv", "Dh"], "cache_keys": ["B", "1", "min(P,W-1)", "Dh"], "cache_keys_post": ["B", "1", "min(P+T,W-1)", "Dh"],
        "c_kv": ["B", "T", "2Dh"], "c_gate": ["B", "T", "2Dh"], "c_chunk_kv": ["B", "N_win×M", "2Dh"], "c_chunk_gate": ["B", "N_win×M", "2Dh"],
        "c_windows_kv": ["B", "N_win", "2M", "Dh"], "c_windows_gate": ["B", "N_win", "2M", "Dh"], "c_compressed": ["B", "N_win", "Dh"], "c_compressed_rope": ["B", "N_win", "Dh"],
        "c_cos": ["B", "N_win", "Rk÷2"], "c_sin": ["B", "N_win", "Rk÷2"], "c_block_bias": ["B", "1", "T", "N_comp"],
        "i_kv": ["B", "T", "2Id"], "i_gate": ["B", "T", "2Id"], "i_chunk_kv": ["B", "N_win_i×M", "2Id"], "i_chunk_gate": ["B", "N_win_i×M", "2Id"],
        "i_windows_kv": ["B", "N_win_i", "2M", "Id"], "i_windows_gate": ["B", "N_win_i", "2M", "Id"], "i_compressed": ["B", "N_win_i", "Id"],
        "i_q": ["B", "T", "Ih", "Id"], "i_scores": ["B", "T", "Tci+N_win_i"], "indexer_topk": ["B", "T", "min(Ki,Tci+N_win_i)"],
        "kv_all": ["B", "1", "L_kv+N_comp", "Dh"], "kv_repeated": ["B", "Nh", "L_kv+N_comp", "Dh"],
        "qk_logits": ["B", "Nh", "T", "L_kv+N_comp"], "qk_masked": ["B", "Nh", "T", "L_kv+N_comp"], "attn_probs": ["B", "Nh", "T", "L_kv+N_comp"],
        "attn_head_out": ["B", "T", "Nh", "Dh"], "attn_rope_inv": ["B", "Nh", "T", "Dh"], "attn_grouped": ["B", "T", "G", "Nh×Dh÷G"],
        "o_a_out": ["B", "T", "G×Ro"], "attn_block_output": ["B", "T", "H"], "attn_combT": ["B", "T", "Nhc", "Nhc"], "attn_expanded": ["B", "T", "Nhc", "H"],
        "ffn_flat": ["B×T", "H"], "router_logits": ["B×T", "E"], "topk_indices": ["B×T", "K"], "topk_weights": ["B×T", "K"],
        "expert_mask": ["E+1", "K", "B×T"], "hit_experts": ["N_hit", "1"], "expert_tokens": ["N_e", "H"],
        "expert_token_idx": ["N_e"], "expert_topk_pos": ["N_e"], "expert_gate_up_e": ["2Es", "H"], "expert_down_e": ["H", "Es"],
        "expert_gu_out": ["N_e", "2Es"], "expert_gate_c": ["N_e", "Es"], "expert_up_c": ["N_e", "Es"], "expert_gated": ["N_e", "Es"],
        "expert_down_raw": ["N_e", "H"], "expert_down_out": ["N_e", "H"], "routed_out": ["B×T", "H"], "routed_view": ["B", "T", "H"],
        "shared_gate": ["B", "T", "Es"], "shared_up": ["B", "T", "Es"], "shared_act": ["B", "T", "Es"], "shared_out": ["B", "T", "H"],
        "mlp_output": ["B", "T", "H"], "ffn_combT": ["B", "T", "Nhc", "Nhc"], "ffn_expanded": ["B", "T", "Nhc", "H"],
        "mask_model": ["B", "1", "T", "L_mask"], "mask_ext": ["B", "1", "T", "L_mask+N_comp"],
        "position_ids": ["B", "T"], "cache_cumlen": [], "cache_entry_count_c": [], "cache_entry_count_i": [],
        "cache_entry_count_c_post": [], "cache_entry_count_i_post": [],
        "cache_compressed_c": ["B", "Tcc", "Dh"], "cache_compressed_c_post": ["B", "N_comp", "Dh"],
        "cache_compressed_i": ["B", "Tci", "Id"], "cache_compressed_i_post": ["B", "Tci+N_win_i", "Id"],
    }
    for site in ("attn", "ffn"):
        shapes[f"{site}_hc_fn"] = ["(2+Nhc)×Nhc", "Nhc×H"]
        shapes[f"{site}_hc_base"] = ["(2+Nhc)×Nhc"]
        shapes[f"{site}_hc_scale"] = ["3"]
    for d, dim in (("c", "Dh"), ("i", "Id")):
        shapes[f"cache_buffer_{d}_kv"] = ["B", "Tbuf", f"2{dim.title()}" if False else ("2Dh" if d == "c" else "2Id")]
        shapes[f"cache_buffer_{d}_gate"] = shapes[f"cache_buffer_{d}_kv"]
        shapes[f"cache_buffer_{d}_kv_post"] = ["B", "Tbuf'", shapes[f"cache_buffer_{d}_kv"][-1]]
        shapes[f"cache_buffer_{d}_gate_post"] = shapes[f"cache_buffer_{d}_kv_post"]
        shapes[f"cache_overlap_{d}_kv"] = ["B", "M", dim]
        shapes[f"cache_overlap_{d}_gate"] = ["B", "M", dim]
        shapes[f"cache_overlap_{d}_kv_post"] = ["B", "M", dim]
        shapes[f"cache_overlap_{d}_gate_post"] = ["B", "M", dim]
    weight_shapes = {
        "q_weight": ["Rq", "H"], "q_a_norm_weight": ["Rq"], "q_b_weight": ["Nh×Dh", "Rq"],
        "kv_weight": ["Dh", "H"], "kv_norm_weight": ["Dh"], "c_kv_weight": ["2Dh", "H"], "c_gate_weight": ["2Dh", "H"],
        "c_position_bias": ["M", "2Dh"], "c_kv_norm_weight": ["Dh"], "i_kv_weight": ["2Id", "H"], "i_gate_weight": ["2Id", "H"],
        "i_position_bias": ["M", "2Id"], "i_kv_norm_weight": ["Id"], "i_q_weight": ["Ih×Id", "Rq"], "i_w_weight": ["Ih", "H"],
        "sinks": ["Nh"], "o_a_weight": ["G×Ro", "Nh×Dh÷G"], "o_b_weight": ["H", "G×Ro"], "router_weight": ["E", "H"],
        "router_bias": ["E"], "expert_gate_up": ["E", "2Es", "H"], "expert_down": ["E", "H", "Es"],
        "shared_gate_weight": ["Es", "H"], "shared_up_weight": ["Es", "H"], "shared_down_weight": ["H", "Es"],
    }
    for tid, sh in weight_shapes.items():
        shapes[tid] = sh
    kind = {tid: "weight" for tid in weight_shapes}
    for tid in ("attn_hc_fn", "ffn_hc_fn", "expert_gate_up_e", "expert_down_e"):
        kind[tid] = "weight"
    for tid in ("expert_token_idx", "expert_topk_pos", "hit_experts", "expert_mask", "topk_indices", "indexer_topk"):
        kind[tid] = "index"
    for tid in ("cache_keys", "cache_keys_post", "kv_full"):
        kind[tid] = "cache"
    for tid in list(shapes):
        if tid.startswith("cache_"):
            kind[tid] = "cache"
    for tid in ("mask_model", "mask_ext", "c_block_bias"):
        kind[tid] = "mask"
    for tid in ("position_ids",):
        kind[tid] = "boundary"
    tensors = [_t(tid, kind.get(tid, "activation"), sh) for tid, sh in shapes.items()]
    param_specs = [
        ("B", None, ()), ("T", None, ()), ("P", 0, ()), ("H", 4096, ("config",)), ("Nh", 64, ("config",)),
        ("Nkv", 1, ("config",)), ("Dh", 512, ("config",)), ("Rq", 1024, ("config",)), ("Rk", 64, ("derived",)),
        ("G", 8, ()), ("Ro", 1024, ()), ("M", 4, ()), ("Ki", 512, ()), ("Ih", 64, ()), ("Id", 128, ()),
        ("W", 128, ()), ("Nhc", 4, ()), ("Sit", 20, ()), ("Eps", 1e-6, ()), ("E", 256, ("config",)),
        ("K", 6, ()), ("Es", 2048, ("config",)), ("Lim", 10.0, ()), ("Scl", 1.5, ()), ("EpsN", 1e-6, ()),
        ("ThetaC", 160000, ("config",)), ("Yf", 16, ()), ("Nwin", None, ()), ("Ncomp", None, ()), ("Tbuf", None, ()), ("Tbufp", None, ()),
    ]
    parameters = []
    for ppid, val, kinds in param_specs:
        refs = [{"path": "x", "locator": "1", "kind": k, "note": ""} for k in kinds]
        parameters.append({"id": ppid, "symbol": ppid, "meaning": ppid, "value": val, "expression": None, "note": "", "source_refs": refs})
    return {"id": pid, "title": pid, "parameters": parameters, "tensors": tensors, "operators": ops,
            "cache_notes": ["synthetic"], "coverage_notes": ["synthetic"]}


def synthetic_document():
    return {
        "schema": "hf_decoder_reading.v1",
        "model": {"model_key": "deepseek_v4_flash_0731_hf", "display_name": "synthetic",
                  "implementation": "transformers", "source_revision": cl.TRANSFORMERS_HEAD,
                  "config_path": cl.CONFIG_REL},
        "scope": {"layer_index": 4, "attention_type": "compressed_sparse_attention", "moe_type": "moe", "analysis_path": "eager", "notes": []},
        "phases": [_phase("prefill"), _phase("decode")],
    }


def codes(issues, severity="error"):
    return [i.code for i in issues if i.severity == severity]


class HappyTests(unittest.TestCase):
    def test_synthetic_passes(self):
        issues = cl.check_layer4(synthetic_document(), WORKSPACE_ROOT)
        self.assertEqual(codes(issues, "error"), [])
        self.assertEqual(codes(issues, "manual"), [])

    def test_scope_and_identity(self):
        doc = synthetic_document()
        doc["scope"]["layer_index"] = 3
        self.assertIn("scope_layer", codes(cl.check_layer4(doc, WORKSPACE_ROOT)))


class Layer4NegativeTests(unittest.TestCase):
    def setUp(self):
        self.doc = synthetic_document()

    def _phase(self):
        return cl.Phase(self.doc["phases"][0])

    def test_combT_rank(self):
        self.doc["phases"][0]["tensors"] = [t if t["id"] != "attn_combT" else {**t, "shape": ["B", "T", "Nhc", "H"]} for t in self.doc["phases"][0]["tensors"]]
        self.assertIn("combT_rank", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_combT_isolated(self):
        ops = self.doc["phases"][0]["operators"]
        for o in ops:
            if o["id"] == "attn_residual_mix":
                o["inputs"] = ["attn_mhc_comb", "streams_in", "attn_expanded"]
        self.assertIn("combT_isolated", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_prob_axis(self):
        self.doc["phases"][0]["tensors"] = [t if t["id"] != "attn_probs" else {**t, "shape": ["B", "Nh", "T", "Nkv"]} for t in self.doc["phases"][0]["tensors"]]
        self.assertIn("prob_axis", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_av_wrong_input(self):
        ops = self.doc["phases"][0]["operators"]
        for o in ops:
            if o["id"] == "av_matmul":
                o["inputs"] = ["q_rope", "kv_repeated"]
        self.assertIn("av_inputs", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_q_norm_missing_weight(self):
        ops = self.doc["phases"][0]["operators"]
        for o in ops:
            if o["id"] == "q_a_norm":
                o["inputs"] = ["q_a"]
        self.assertIn("q_norm_weight", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_indexer_raw_q(self):
        ops = self.doc["phases"][0]["operators"]
        for o in ops:
            if o["id"] == "i_q_proj":
                o["inputs"] = ["q_a", "cos_q", "sin_q"]
        self.assertIn("indexer_query_norm", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_expert_slice_shape(self):
        self.doc["phases"][0]["tensors"] = [t if t["id"] != "expert_down_e" else {**t, "shape": ["Es", "H"]} for t in self.doc["phases"][0]["tensors"]]
        self.assertIn("expert_slice_shape", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_expert_input_linkage(self):
        ops = self.doc["phases"][0]["operators"]
        for o in ops:
            if o["id"] == "moe_expert_gu":
                o["inputs"] = ["expert_gu_out"]
        self.assertIn("expert_input_linkage", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_expert_weight_apply_missing_topk(self):
        ops = self.doc["phases"][0]["operators"]
        for o in ops:
            if o["id"] == "moe_expert_weight":
                o["inputs"] = ["expert_down_raw", "expert_token_idx", "expert_topk_pos"]
        self.assertIn("expert_weight_apply", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_block_bias_order(self):
        ops = self.doc["phases"][0]["operators"]
        bias = next(o for o in ops if o["id"] == "c_block_bias")
        ops.remove(bias)
        idx = next(i for i, o in enumerate(ops) if o["id"] == "i_topk")
        ops.insert(idx, bias)
        self.assertIn("block_bias_order", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_kv_persist_len(self):
        self.doc["phases"][0]["tensors"] = [t if t["id"] != "cache_keys_post" else {**t, "shape": ["B", "1", "min(P,W-1)", "Dh"]} for t in self.doc["phases"][0]["tensors"]]
        self.assertIn("kv_persist_len", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_overlap_shape(self):
        self.doc["phases"][0]["tensors"] = [t if t["id"] != "cache_overlap_c_kv" else {**t, "shape": ["B", "Dh"]} for t in self.doc["phases"][0]["tensors"]]
        self.assertIn("overlap_shape", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_param_value(self):
        for p in self.doc["phases"][0]["parameters"]:
            if p["id"] == "Es":
                p["value"] = 4096
        self.assertIn("param_value", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_unresolved_declaration(self):
        self.doc["phases"][0]["tensors"].append(_t("mystery_activation", "activation", ["B", "T", "H"]))
        self.assertIn("unresolved_declaration", codes(cl.check_layer4(self.doc, WORKSPACE_ROOT)))

    def test_redundant_declaration_is_adjudicated(self):
        doc = synthetic_document()
        doc["phases"][0]["tensors"].append(_t("probs", "activation", ["B", "Nh", "T", "L_kv+N_comp+1"]))
        issues = cl.check_layer4(doc, WORKSPACE_ROOT)
        self.assertEqual(codes(issues, "error"), [])
        self.assertIn("redundant_declaration", codes(issues, "info"))


class HtmlTests(unittest.TestCase):
    def test_clean(self):
        html = '<section class="phase" id="phase-prefill"></section><section class="phase" id="phase-decode"></section>'
        self.assertEqual(ch.check_html(html, synthetic_document()), [])

    def test_dangling_href(self):
        html = '<section class="phase" id="phase-prefill"></section><section class="phase" id="phase-decode"></section><a href="#gone">x</a>'
        self.assertIn("html_dangling_href", codes(ch.check_html(html, synthetic_document())))

    def test_phase_order(self):
        html = '<section class="phase" id="phase-decode"></section><section class="phase" id="phase-prefill"></section>'
        self.assertIn("html_phase_order", codes(ch.check_html(html, synthetic_document())))

    def test_blackbox_marker(self):
        doc = synthetic_document()
        doc["phases"][0]["operators"][0]["analysis_status"] = "internal_unanalyzed"
        html = '<section class="phase" id="phase-prefill"></section><section class="phase" id="phase-decode"></section>'
        self.assertIn("html_blackbox_marker", codes(ch.check_html(html, doc)))


class WorkbookTests(unittest.TestCase):
    def _render(self, doc, path):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for pid, sheet in (("prefill", "Prefill"), ("decode", "Decode")):
            ws = wb.create_sheet(sheet)
            ph = next(p for p in doc["phases"] if p["id"] == pid)
            ws.append(["参数id", "符号", "含义", "值", "值类型", "表达式", "来源类型", "来源定位", "说明"])
            for pp in ph["parameters"]:
                ws.append([pp["id"], pp["symbol"], pp["meaning"], str(pp["value"]), "type", "", "", "", ""])
            ws.append([None] * 9)
            ws.append(["算子区（synthetic）"])
            ws.append(["黑盒标记", "算子id", "模块路径", "名称", "算子类型", "输入 名称[shape]", "输出 名称[shape]", "权重 名称[shape]", "参数引用 本表id"])
            tmap = {t["id"]: t for t in ph["tensors"]}
            def render(ids):
                return ", ".join(f"{i}[{', '.join(tmap[i]['shape'])}]" for i in ids)
            for o in ph["operators"]:
                ws.append(["已展开", o["id"], o["module_path"], o["name"], o["op"], render(o["inputs"]), render(o["outputs"]), render(o["weights"]), ", ".join(o["parameter_refs"] or [ph["parameters"][0]["id"]])])
        wb.save(path)

    def test_clean(self):
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "wb.xlsx"
            self._render(synthetic_document(), path)
            self.assertEqual(codes(cx.check_xlsx(path, synthetic_document())), [])

    def test_wrong_sheet_order(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "wb.xlsx"
            wb = openpyxl.Workbook(); wb.remove(wb.active); wb.create_sheet("Decode"); wb.create_sheet("Prefill"); wb.save(path)
            self.assertIn("xlsx_sheets", codes(cx.check_xlsx(path, synthetic_document())))

    def test_dangling_param_link(self):
        doc = synthetic_document()
        doc["phases"][0]["operators"][0]["parameter_refs"] = ["NOPE"]
        # render maps empty params to first param; force explicit bad token
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "wb.xlsx"
            self._render(doc, path)
            # rewrite the first operator row's 参数引用 cell to the bad token
            wb = openpyxl.load_workbook(path)
            ws = wb["Prefill"]
            hdr_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "算子id")
            ws.cell(hdr_row + 1, 9, "NOPE")
            wb.save(path)
            self.assertIn("xlsx_param_link", codes(cx.check_xlsx(path, doc)))


if __name__ == "__main__":
    unittest.main()
