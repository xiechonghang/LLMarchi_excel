"""Independent read-back checks for the two-sheet xlsx (no native Excel)."""

from __future__ import annotations

import re
from pathlib import Path

from check_layer4 import Issue, _err, _warn, _manual  # type: ignore


def _find_header(ws, token):
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and token in v:
                return r, c
    return None, None


def _cols(ws, header_row):
    return {ws.cell(header_row, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(header_row, c).value}


def check_xlsx(path: Path, document: dict):
    issues: list[Issue] = []
    try:
        import openpyxl
    except Exception as exc:  # pragma: no cover
        return [_manual("xlsx_openpyxl", f"openpyxl unavailable: {exc}")]
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:
        return [_err("xlsx_unreadable", f"cannot read workbook: {exc}")]
    if wb.sheetnames != ["Prefill", "Decode"]:
        issues.append(_err("xlsx_sheets", f"sheets must be ['Prefill','Decode'], got {wb.sheetnames}"))
    phases = {p["id"]: p for p in document.get("phases", [])}
    for pid, sheet in (("prefill", "Prefill"), ("decode", "Decode")):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        pr, pc = _find_header(ws, "参数id")
        orow, oidc = _find_header(ws, "算子id")
        if pr is None or orow is None:
            issues.append(_err("xlsx_regions", f"{sheet}: missing 参数id / 算子id header", pid))
            continue
        doc = phases.get(pid, {})
        # param ids
        param_ids = []
        r = pr + 1
        while r <= ws.max_row and ws.cell(r, pc).value not in (None, ""):
            param_ids.append(str(ws.cell(r, pc).value).strip())
            r += 1
        doc_params = [p["id"] for p in doc.get("parameters", [])]
        if set(param_ids) != set(doc_params):
            issues.append(_err("xlsx_param_ids", f"{sheet}: param id set differs from JSON (sheet {len(param_ids)} vs doc {len(doc_params)})", pid))
        param_set = set(param_ids)
        # operator rows
        ocols = _cols(ws, orow)
        def col(tok):
            return next((c for k, c in ocols.items() if isinstance(k, str) and tok in k), None)
        cin, cout, cw, cref, cname = col("输入"), col("输出"), col("权重"), col("参数引用"), col("名称")
        # no repetition-count column anywhere in the header
        for k in ocols:
            if isinstance(k, str) and ("重复次数" in k or "repetition" in k.lower() or "repeat_count" in k.lower()):
                issues.append(_err("xlsx_repeat_column", f"{sheet}: forbidden repetition-count column {k!r}", pid))
        tmap = {t["id"]: t for t in doc.get("tensors", [])}
        doc_ops = {o["id"]: o for o in doc.get("operators", [])}
        sheet_ops = {}
        r = orow + 1
        while r <= ws.max_row and ws.cell(r, oidc).value not in (None, ""):
            sheet_ops[str(ws.cell(r, oidc).value).strip()] = r
            r += 1
        if set(sheet_ops) != set(doc_ops):
            issues.append(_err("xlsx_op_ids", f"{sheet}: operator id set differs from JSON (sheet {len(sheet_ops)} vs doc {len(doc_ops)})", pid))
        shape_mismatch = name_mismatch = link_bad = 0
        for oid, rr in sheet_ops.items():
            o = doc_ops.get(oid)
            if not o:
                continue
            if cname and ws.cell(rr, cname).value is not None and str(ws.cell(rr, cname).value).strip() != str(o.get("name", "")).strip():
                name_mismatch += 1
            if cref and ws.cell(rr, cref).value:
                for tok in re.split(r"[,;/、\s]+", str(ws.cell(rr, cref).value).strip()):
                    if tok and tok not in param_set:
                        link_bad += 1
            for ids, ci in ((o.get("inputs"), cin), (o.get("outputs"), cout), (o.get("weights"), cw)):
                if ci is None or not ids:
                    continue
                cell = str(ws.cell(rr, ci).value or "")
                for tid in ids:
                    sh = tmap.get(tid, {}).get("shape")
                    if sh is None:
                        continue
                    if "[" + ", ".join(sh) + "]" not in cell:
                        shape_mismatch += 1
        if name_mismatch:
            issues.append(_err("xlsx_op_name", f"{sheet}: {name_mismatch} operator name(s) differ from JSON", pid))
        if link_bad:
            issues.append(_err("xlsx_param_link", f"{sheet}: {link_bad} parameter reference token(s) do not resolve in this sheet", pid))
        if shape_mismatch:
            issues.append(_err("xlsx_shape_mismatch", f"{sheet}: {shape_mismatch} tensor shape(s) from JSON not found in operator rows", pid))
    return issues


__all__ = ["check_xlsx"]
